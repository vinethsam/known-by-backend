from __future__ import annotations

import csv
import io
import unicodedata
import zipfile
from dataclasses import dataclass, field
from pathlib import PurePath
from typing import Any
from xml.etree.ElementTree import ParseError, XMLParser, iterparse

from openpyxl import load_workbook
from openpyxl.utils.cell import range_boundaries
from openpyxl.utils.exceptions import InvalidFileException
from pydantic import ValidationError

from app.config import Settings
from app.input_validation import (
    MAX_INPUT_CELL_CHARS,
    MAX_INPUT_FILENAME_CHARS,
    MAX_INPUT_HEADER_CHARS,
    sanitize_input_text,
    validate_input_text,
)
from app.schemas import PersonSeed
from app.seed_context import _OBVIOUS_EMAIL, COLUMN_ALIASES, _classify_header, safe_context_value
from app.seed_context import normalize_header as normalize_header


@dataclass(frozen=True)
class _ColumnInference:
    name_column: str
    seed_columns: dict[str, str]


@dataclass(frozen=True)
class BatchInput:
    seeds: list[PersonSeed | None]
    rows: list[dict[str, Any]]
    columns: list[str]
    row_errors: list[str | None] = field(default_factory=list)


class BatchParseError(ValueError):
    pass


def parse_batch(
    content: bytes,
    filename: str,
    name_column: str | None,
    settings: Settings,
) -> BatchInput:
    if len(content) > settings.MAX_UPLOAD_BYTES:
        raise BatchParseError("Batch upload exceeds maximum size")
    _validate_text(filename, MAX_INPUT_FILENAME_CHARS)
    if name_column is not None:
        _validate_text(name_column, MAX_INPUT_HEADER_CHARS)
    suffix = PurePath(filename).suffix.lower()
    if suffix == ".xlsx":
        rows = _read_xlsx(content, settings)
    elif suffix == ".tsv":
        rows = _read_delimited(content, "\t", settings)
    elif suffix == ".csv" or not suffix:
        rows = _read_delimited(content, ",", settings)
    else:
        raise BatchParseError("Unsupported batch file type")
    return _rows_to_batch(rows, name_column, settings)


def _read_delimited(content: bytes, delimiter: str, settings: Settings) -> list[list[str]]:
    try:
        text = content.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise BatchParseError("Batch file must be UTF-8 text") from exc
    reader = csv.reader(io.StringIO(text, newline=""), delimiter=delimiter, strict=True)
    rows: list[list[str]] = []
    try:
        for row_number, row in enumerate(reader, start=1):
            if row_number > settings.MAX_BATCH_ROWS + 1:
                raise BatchParseError(
                    f"Batch file contains too many rows; maximum {settings.MAX_BATCH_ROWS} data records"
                )
            if len(row) > settings.MAX_INPUT_COLUMNS:
                raise BatchParseError("Batch file contains too many columns")
            rows.append([_stringify_cell(value) for value in row])
    except csv.Error as exc:
        raise BatchParseError("CSV is malformed or contains an oversized cell") from exc
    return rows


def _read_xlsx(content: bytes, settings: Settings) -> list[list[str]]:
    _validate_xlsx_archive(content, settings)
    try:
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=False, keep_links=False)
    except (InvalidFileException, OSError, KeyError, ValueError, ParseError, zipfile.BadZipFile) as exc:
        raise BatchParseError("XLSX file is corrupted") from exc
    try:
        sheet = workbook.active
        if sheet is None:
            raise BatchParseError("XLSX file has no active worksheet")
        _validate_xlsx_dimensions(sheet, settings)
        # Validate actual cell coordinates as well as the declared worksheet range.
        # A forged small dimension must not silently hide rows/columns from a batch.
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            with archive.open(sheet._worksheet_path) as xml:
                for _, element in iterparse(xml, events=("end",)):
                    if element.tag.endswith("}c") and element.get("r"):
                        _, _, column, row_index = range_boundaries(element.get("r"))
                        if row_index > settings.MAX_BATCH_ROWS + 1:
                            raise BatchParseError(
                                f"XLSX cells exceed configured row limit; maximum {settings.MAX_BATCH_ROWS} data records"
                            )
                        if column > settings.MAX_INPUT_COLUMNS:
                            raise BatchParseError("XLSX cells exceed configured columns")
                    element.clear()
        sheet.reset_dimensions()
        rows: list[list[str]] = []
        for row_number, row in enumerate(
            sheet.iter_rows(),
            start=1,
        ):
            if row_number > settings.MAX_BATCH_ROWS + 1:
                raise BatchParseError(
                    f"Batch file contains too many rows; maximum {settings.MAX_BATCH_ROWS} data records"
                )
            values = [_stringify_cell(cell.value) for cell in row]
            if len(values) > settings.MAX_INPUT_COLUMNS:
                raise BatchParseError("Batch file contains too many columns")
            rows.append(values)
        return rows
    except (ParseError, KeyError, TypeError, ValueError, OSError, zipfile.BadZipFile) as exc:
        if isinstance(exc, BatchParseError):
            raise
        raise BatchParseError("XLSX file is corrupted") from exc
    finally:
        workbook.close()


def _validate_xlsx_archive(content: bytes, settings: Settings) -> None:
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            entries = archive.infolist()
            if len(entries) > 500:
                raise BatchParseError("XLSX file contains too many internal parts")
            if len({info.filename for info in entries}) != len(entries):
                raise BatchParseError("XLSX file contains duplicate internal parts")
            total = 0
            for info in entries:
                if info.flag_bits & 1:
                    raise BatchParseError("Encrypted XLSX files are not supported")
                total += info.file_size
                if total > settings.MAX_XLSX_UNCOMPRESSED_BYTES:
                    raise BatchParseError("XLSX file expands beyond the maximum allowed size")
                if info.compress_size and info.file_size / info.compress_size > 1000:
                    raise BatchParseError("XLSX compression ratio is too large")
            # Parse without building trees before openpyxl can expand entities or
            # allocate shared strings. The parser recognizes DTDs in UTF-16 too.
            for info in entries:
                if info.filename.lower().endswith((".xml", ".rels")):
                    parser = XMLParser(target=_BoundedXMLTarget())
                    with archive.open(info) as part:
                        while chunk := part.read(65536):
                            parser.feed(chunk)
                    parser.close()
    except (zipfile.BadZipFile, OSError, ValueError, ParseError, RuntimeError, NotImplementedError) as exc:
        if isinstance(exc, BatchParseError):
            raise
        raise BatchParseError("XLSX file is corrupted") from exc


class _BoundedXMLTarget:
    def start(self, tag, attributes):
        pass

    def end(self, tag):
        pass

    def data(self, value):
        pass

    def close(self):
        pass

    def doctype(self, name, public_id, system_id):
        raise BatchParseError("XLSX XML document types and entities are not supported")


def _rows_to_batch(rows: list[list[str]], name_column: str | None, settings: Settings) -> BatchInput:
    if not rows:
        raise BatchParseError("Batch file is empty")
    header = []
    for value in rows[0]:
        try:
            _validate_text(value, MAX_INPUT_HEADER_CHARS, allow_controls=True)
            _validate_text(sanitize_input_text(value), MAX_INPUT_HEADER_CHARS)
        except ValueError as exc:
            raise BatchParseError(str(exc)) from exc
        header.append(value)
    _validate_headers(header, settings)

    prepared_rows: list[tuple[int, list[str]]] = []
    for zero_based, raw_row in enumerate(rows[1:], start=2):
        if len(raw_row) > len(header):
            raise BatchParseError(f"Row {zero_based} contains data beyond the header")
        values = list(raw_row)
        if len(values) < len(header):
            values.extend([""] * (len(header) - len(values)))
        if len(prepared_rows) >= settings.MAX_BATCH_ROWS:
            raise BatchParseError(
                f"Batch file contains too many rows; maximum {settings.MAX_BATCH_ROWS} data records"
            )
        prepared_rows.append((zero_based, values))

    inference = _infer_columns(header, [values for _, values in prepared_rows], name_column)
    output_rows: list[dict[str, Any]] = []
    seeds: list[PersonSeed | None] = []
    row_errors: list[str | None] = []
    for row_number, values in prepared_rows:
        row = {column: values[index] for index, column in enumerate(header)}
        output_rows.append(row)
        try:
            for value in values:
                _validate_text(value, MAX_INPUT_CELL_CHARS, multiline=True, allow_controls=True)
            name = sanitize_input_text(row[inference.name_column]).strip()
            if not name or _obviously_not_person_name(name):
                seeds.append(None)
                row_errors.append("INVALID_PERSON_NAME")
                continue
            seed = _seed_from_row(row, name, inference.seed_columns, name_column=inference.name_column)
        except ValueError:
            seeds.append(None)
            row_errors.append("INVALID_PERSON_SEED")
        else:
            seeds.append(seed)
            row_errors.append(None)
    if not seeds:
        raise BatchParseError("Batch file contains no people")
    return BatchInput(seeds=seeds, rows=output_rows, columns=header, row_errors=row_errors)


def _validate_headers(header: list[str], settings: Settings) -> None:
    if not header:
        raise BatchParseError("Batch file is missing a header row")
    if len(header) > settings.MAX_INPUT_COLUMNS:
        raise BatchParseError("Batch file contains too many columns")
    if any(not sanitize_input_text(column).strip() for column in header):
        raise BatchParseError("Batch file contains a blank header")
    normalized = [sanitize_input_text(column).strip().casefold() for column in header]
    if len(set(normalized)) != len(normalized):
        raise BatchParseError("Batch file contains duplicate headers")


def _infer_columns(
    header: list[str], data_rows: list[list[str]], name_column: str | None
) -> _ColumnInference:
    column_indexes = {column: index for index, column in enumerate(header)}
    matches = [match for column in header if (match := _classify_header(column)) is not None]
    if name_column:
        explicit = [
            (index, column)
            for index, column in enumerate(header)
            if sanitize_input_text(column) == unicodedata.normalize("NFC", name_column)
        ]
        if not explicit:
            raise BatchParseError("Selected name column is not present")
        _, selected_name = explicit[0]
    else:
        name_candidates = [match for match in matches if match.field == "full_name"]
        usable = [
            match for match in name_candidates if _usable_name_column(column_indexes[match.column], data_rows)
        ]
        # Prefer columns with real names, as before. If every recognized name is
        # invalid, still retain those records and report their individual errors.
        name_candidates = usable or name_candidates
        if not name_candidates:
            observed = ", ".join(repr(column) for column in header)
            raise BatchParseError(
                "Batch file is missing a name column; no usable person-name candidate "
                f"was inferred. Observed headers: {observed}"
            )
        best_score = max(match.score for match in name_candidates)
        strongest = [match for match in name_candidates if match.score == best_score]
        if len(strongest) > 1:
            candidates = ", ".join(repr(match.column) for match in strongest)
            raise BatchParseError(f"Batch file has ambiguous name columns: {candidates}")
        selected_name = strongest[0].column

    seed_columns: dict[str, str] = {}
    for field_name in COLUMN_ALIASES:
        if field_name == "full_name":
            continue
        candidates = [
            match for match in matches if match.field == field_name and match.column != selected_name
        ]
        if not candidates:
            continue
        best_score = max(match.score for match in candidates)
        strongest = [match for match in candidates if match.score == best_score]
        # Equally strong optional columns are retained in the original row but are
        # not guessed into a single-valued PersonSeed field.
        if len(strongest) == 1:
            seed_columns[field_name] = strongest[0].column
    return _ColumnInference(name_column=selected_name, seed_columns=seed_columns)


def _usable_name_column(index: int, data_rows: list[list[str]]) -> bool:
    if not data_rows:
        return True
    for row in data_rows:
        if index >= len(row):
            continue
        try:
            value = sanitize_input_text(row[index]).strip()
        except ValueError:
            continue
        if value and not _obviously_not_person_name(value):
            return True
    return False


def _obviously_not_person_name(value: str) -> bool:
    normalized = value.strip().casefold()
    if normalized in {"-", "n/a", "na", "none", "null", "unknown"}:
        return True
    if normalized.startswith(("http://", "https://", "www.")) or "://" in normalized:
        return True
    if _OBVIOUS_EMAIL.fullmatch(normalized):
        return True
    compact = "".join(character for character in normalized if character.isalnum())
    return not compact or compact.isdecimal() or not any(character.isalpha() for character in compact)


def _seed_from_row(
    row: dict[str, str], full_name: str, seed_columns: dict[str, str], *, name_column: str | None = None
) -> PersonSeed:
    processed = {
        field_name: sanitize_input_text(row[column]).strip() for field_name, column in seed_columns.items()
    }
    values = {
        field_name: value
        for field_name, value in processed.items()
        if value
        and (field_name == "profile_link" or safe_context_value(value))
        and len(value)
        <= {"country": 100, "country_of_origin": 100, "program_year": 20, "profile_link": 2048}.get(
            field_name, 200
        )
    }
    semantic_fields = {
        "current_employer",
        "alumni_organisation",
        "government_office",
        "nationality",
        "citizenship",
        "birthplace",
    }
    payload: dict[str, Any] = {
        "full_name": full_name,
        "input_metadata": dict(row),
        "input_field_map": {
            column: match.field for column in row if (match := _classify_header(column)) is not None
        },
        **{field_name: value for field_name, value in values.items() if field_name not in semantic_fields},
    }
    if name_column:
        payload["input_field_map"][name_column] = "full_name"

    # Header semantics are retained inside the existing PersonSeed context. These
    # values guide research only; they never create evidence claims or provenance.
    known_attributes = {
        field_name: values[field_name]
        for field_name in (
            "current_employer",
            "government_office",
            "alumni_organisation",
            "historical_employer",
        )
        if field_name in values
    }

    # An explicit current-employer header is the strongest current affiliation
    # hypothesis. A government-office header is an institutional identity anchor
    # when no ordinary/current organisation was supplied. Alumni and former
    # employers deliberately remain supporting context only.
    generic_organisation = values.get("organisation")
    selected_organisation = (
        values.get("current_employer") or generic_organisation or values.get("government_office")
    )
    if selected_organisation:
        payload["organisation"] = selected_organisation
    if generic_organisation and generic_organisation != selected_organisation:
        known_attributes["organisation_affiliation"] = generic_organisation
    if known_attributes:
        payload["known_attributes"] = known_attributes
    try:
        return PersonSeed.model_validate(payload)
    except ValidationError as exc:
        raise BatchParseError(f"Invalid person seed: {exc.errors()[0]['msg']}") from exc


def _stringify_cell(value: Any) -> str:
    if value is None:
        return ""
    value = str(value)
    # Keep file-level size safeguards, but defer character validity to each row.
    if len(value) > MAX_INPUT_CELL_CHARS:
        _validate_text(value, MAX_INPUT_CELL_CHARS, multiline=True, allow_controls=True)
    return value


def _validate_text(
    value: str, max_length: int, *, multiline: bool = False, allow_controls: bool = False
) -> None:
    try:
        validate_input_text(value, max_length=max_length, multiline=multiline, allow_controls=allow_controls)
    except ValueError as exc:
        raise BatchParseError(str(exc)) from exc


def _validate_xlsx_dimensions(sheet, settings: Settings) -> None:
    try:
        minimum_column, minimum_row, maximum_column, maximum_row = range_boundaries(
            sheet.calculate_dimension()
        )
    except ValueError:
        # Unsized sheets are checked through actual XML coordinates above.
        return
    if minimum_column < 1 or minimum_row < 1:
        raise BatchParseError("XLSX file has invalid dimensions")
    if maximum_column > settings.MAX_INPUT_COLUMNS:
        raise BatchParseError("Batch file contains too many columns")
    if maximum_row > settings.MAX_BATCH_ROWS + 1:
        raise BatchParseError(
            f"Batch file contains too many rows; maximum {settings.MAX_BATCH_ROWS} data records"
        )
