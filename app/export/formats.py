from __future__ import annotations

import csv
import io
import json
import re
from collections.abc import Iterable, Iterator
from typing import Any, Literal

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app.config import DEFAULT_REVIEW_THRESHOLD
from app.export.schema import RESEARCH_FIELDS, standard_columns
from app.research.normalisation import finalize_display
from app.research.seed_provenance import project_stored_seed, supplied_field_decisions
from app.retrieval.urls import source_policy_allows, source_record_allowed
from app.schemas import JobResults, ProfileField

FORMULA_PREFIXES = ("=", "+", "-", "@", "\t", "\r")
ProvenanceMode = Literal["none", "field"]
_JSON_FORMULA_ESCAPES = str.maketrans({char: f"\\u{ord(char):04x}" for char in "=+-@"})
MAX_EXPORT_ROWS = 500
COMBINED_EDUCATION_REASON = "EDUCATION_RECORDS_COMBINED_FOR_EXPORT_LIMIT"


class ExportLimitError(ValueError):
    """The full result cannot fit the standardized artifact limits."""


class ExportCellLimitError(ExportLimitError):
    """A standardized cell cannot be represented losslessly in both formats."""


def _provenance_json(value, *, compact=False):
    payload = json.dumps(value, ensure_ascii=False, separators=(",", ":") if compact else None)
    # Keep JSON fragments safe as spreadsheet cells without changing decoded input.
    # Otherwise a split inside a string can expose a formula prefix, causing the
    # ordinary cell sanitizer to insert an apostrophe into the original value.
    return re.sub(
        r'"(?:[^"\\]|\\.)*"',
        lambda match: match.group().translate(_JSON_FORMULA_ESCAPES),
        payload,
    )


def export_results(
    results: JobResults,
    columns: list[str],
    format: str,
    *,
    provenance: ProvenanceMode = "none",
    low_confidence_threshold: float = DEFAULT_REVIEW_THRESHOLD,
) -> bytes:
    if provenance not in {"none", "field"}:
        raise ValueError("Unsupported provenance mode")
    rows, output_columns = _flatten_results(results, columns, provenance=provenance)
    rows = _validated_rows(rows, output_columns)
    normalized = format.lower()
    if normalized == "csv":
        return _export_csv(rows, output_columns)
    if normalized == "xlsx":
        return _export_xlsx(
            rows,
            output_columns,
            field_columns=_field_column_pairs(),
            low_confidence_threshold=low_confidence_threshold,
        )
    raise ValueError("Unsupported export format")


def _flatten_results(
    results: JobResults,
    columns: list[str],
    *,
    provenance: ProvenanceMode = "none",
) -> tuple[Iterator[dict[str, Any]], list[str]]:
    # Input columns and the legacy provenance switch no longer alter normal files.
    # Full raw input remains in storage/API; the artifact contains standardized data.
    if len(results.people) > MAX_EXPORT_ROWS:
        raise ExportLimitError(
            "Export supports at most 500 people; the complete research remains available in record details"
        )
    expanded = sum(
        max(1, len(person.result.profile.records)) if person.result else 1 for person in results.people
    )
    return _export_rows(results, compact=expanded > MAX_EXPORT_ROWS), list(standard_columns())


def _export_rows(results: JobResults, *, compact: bool) -> Iterator[dict[str, Any]]:
    if not compact:
        yield from _iter_rows(results)
        return
    for person in results.people:
        # Only one person's education rows are buffered. Durable API/history
        # records are never replaced by this explicitly marked export projection.
        rows = list(_iter_rows(results.model_copy(update={"people": [person]})))
        yield _combine_education_rows(rows) if len(rows) > 1 else rows[0]


def _combine_education_rows(rows: list[dict[str, Any]]) -> dict[str, Any]:
    combined = rows[0].copy()
    combined["_review_fields"] = set().union(*(row["_review_fields"] for row in rows))
    for field in RESEARCH_FIELDS:
        key = field.value
        decisions = [row for row in rows if row.get(key)]
        values = {row.get(key, "") for row in rows}
        numbered = bool(decisions) and (
            key in {"university_name", "degree_type", "subject"} or len(values) > 1
        )
        if numbered:
            combined[key] = "\n".join(f"{index}. {row.get(key) or '—'}" for index, row in enumerate(rows, 1))
            combined["_review_fields"].add(key)
        scores = [row[f"{key}_confidence"] for row in decisions]
        combined[f"{key}_confidence"] = min(scores) if scores else 0
        combined[f"{key}_source_urls"] = _compact_urls(
            url for row in rows for url in row.get(f"{key}_source_urls", "").split("; ")
        )
        supplied = list(
            dict.fromkeys(
                value for row in rows for value in json.loads(row.get(f"{key}_supplied_values", "[]"))
            )
        )
        combined[f"{key}_supplied_values"] = _provenance_json(supplied)
        combined[f"{key}_input_confidence"] = 100 if supplied else ""
        states = {row[f"{key}_provenance_status"] for row in decisions}
        combined[f"{key}_provenance_status"] = (
            next(iter(states)) if len(states) == 1 else "mixed" if states else ""
        )
        reasons = list(
            dict.fromkeys(
                reason
                for row in rows
                for reason in row.get(f"{key}_review_reasons", "").split("; ")
                if reason
            )
        )
        if numbered:
            reasons.append(COMBINED_EDUCATION_REASON)
        combined[f"{key}_review_required"] = (
            any(row.get(f"{key}_review_required") for row in rows) or key in combined["_review_fields"]
        )
        combined[f"{key}_review_reasons"] = "; ".join(dict.fromkeys(reasons))
    for key in ("profile_confidence", "coverage", "legacy_coverage"):
        combined[key] = min(row[key] for row in rows)
    combined["review_required"] = True
    combined["record_status"] = (
        combined["status"] if combined["status"] in {"failed", "cancelled"} else "review_required"
    )
    reasons = [reason for row in rows for reason in row.get("review_reasons", "").split("; ") if reason]
    combined["review_reasons"] = "; ".join(dict.fromkeys([*reasons, COMBINED_EDUCATION_REASON]))
    return combined


def _iter_rows(results: JobResults) -> Iterator[dict[str, Any]]:
    for person in results.people:
        if person.result:
            profile = finalize_display(
                project_stored_seed(
                    person.result.profile,
                    {
                        "full_name": person.result.profile.input_name,
                        "input_metadata": person.original_row,
                    },
                )
            )
            records = profile.records or [profile]
            source_urls = _compact_source_urls(person.result)
        else:
            profile = None
            records = [None]
            source_urls = ""
        for record in records:
            fields = (
                record.fields
                if record
                else (
                    person.supplied_fields
                    or supplied_field_decisions({"input_metadata": person.original_row})
                )
            )
            row = {
                "status": person.status.value,
                "record_status": record.status.value if record else person.status.value,
                "error_code": person.error_code or "",
                "profile_confidence": record.profile_confidence if record else 0,
                "coverage": record.coverage if record else 0,
                "legacy_coverage": record.legacy_coverage if record else 0,
                "review_required": record.review_required if record else False,
                "review_reasons": "; ".join(
                    getattr(record, "review_reason_codes", [])
                    or sorted(
                        {reason for decision in fields.values() for reason in decision.review_reason_codes}
                    )
                ),
                "source_urls": source_urls,
                "_review_fields": set(),
            }
            for field in RESEARCH_FIELDS:
                decision = fields.get(field)
                if not decision:
                    continue
                value = decision.value
                # Supplied links are input data, never asserted research sources.
                blocked_link = bool(
                    field == ProfileField.profile_link and value and not source_policy_allows(value)
                )
                if blocked_link:
                    value = None
                row[field.value] = value or ""
                row[f"{field.value}_confidence"] = (
                    0
                    if blocked_link
                    else decision.research_confidence
                    if decision.research_confidence is not None
                    else decision.confidence
                )
                row[f"{field.value}_input_confidence"] = (
                    decision.input_confidence if decision.input_confidence is not None else ""
                )
                row[f"{field.value}_provenance_status"] = decision.provenance_status or (
                    "researched" if value else ""
                )
                row[f"{field.value}_supplied_values"] = _provenance_json(decision.supplied_values)
                row[f"{field.value}_source_urls"] = _compact_urls(decision.sources)
                row[f"{field.value}_review_required"] = decision.review_required
                row[f"{field.value}_review_reasons"] = "; ".join(decision.review_reason_codes)
                if decision.review_required or decision.provenance_status in {"conflicting", "superseded"}:
                    row["_review_fields"].add(field.value)
            yield row


def _validated_rows(rows, columns):
    for row in rows:
        if any(
            isinstance(row.get(column), str) and len(_safe_cell(row[column])) > 32767 for column in columns
        ):
            raise ExportCellLimitError(
                "A result cell exceeds the spreadsheet limit of 32767 characters; "
                "original supplied values remain available in record details."
            )
        yield row


def _field_column_pairs() -> list[tuple[str, str]]:
    return [(field.value, f"{field.value}_confidence") for field in RESEARCH_FIELDS]


def _compact_source_urls(result) -> str:
    return _compact_urls(
        (
            source.canonical_url or source.final_url or source.requested_url
            for source in result.sources
            if source_record_allowed(source)
        ),
        limit=10,
    )


def _compact_urls(values, *, limit: int | None = None) -> str:
    urls: list[str] = []
    seen: set[str] = set()
    for value in values:
        url = str(value)
        if not url or url in seen:
            continue
        seen.add(url)
        if source_policy_allows(url):
            urls.append(url)
            if limit is not None and limit > 0 and len(urls) >= limit:
                break
    return "; ".join(urls if limit is None else urls[:limit])


def _safe_cell(value: Any) -> Any:
    if isinstance(value, str) and value.lstrip().startswith(FORMULA_PREFIXES):
        return "'" + value
    return value


def _export_csv(rows: Iterable[dict[str, Any]], columns: list[str]) -> bytes:
    output = io.StringIO(newline="")
    writer = csv.DictWriter(output, fieldnames=columns, extrasaction="ignore")
    writer.writerow({column: _safe_cell(column) for column in columns})
    for row in rows:
        writer.writerow({column: _safe_cell(row.get(column, "")) for column in columns})
    return output.getvalue().encode("utf-8-sig")


def _export_xlsx(
    rows: Iterable[dict[str, Any]],
    columns: list[str],
    *,
    field_columns: list[tuple[str, str]] = (),
    low_confidence_threshold: float = DEFAULT_REVIEW_THRESHOLD,
) -> bytes:
    if not 0 <= low_confidence_threshold <= 100:
        raise ValueError("Low-confidence threshold must be between 0 and 100")
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Results"
    header_fill = PatternFill(fill_type="solid", fgColor="D9EAF7")
    header_font = Font(bold=True)
    for index, column in enumerate(columns, start=1):
        cell = sheet.cell(row=1, column=index, value=_xlsx_cell(column))
        cell.font = header_font
        cell.fill = header_fill
    pale_yellow = PatternFill(fill_type="solid", fgColor="FFF2CC")
    column_indexes = {column: index for index, column in enumerate(columns, start=1)}
    widths = [len(str(column)) for column in columns]
    for row_index, row in enumerate(rows, start=2):
        for column_index, column in enumerate(columns, start=1):
            value = row.get(column, "")
            sheet.cell(row=row_index, column=column_index, value=_xlsx_cell(value))
            if isinstance(value, str) and "\n" in value:
                sheet.cell(row=row_index, column=column_index).alignment = Alignment(
                    wrap_text=True, vertical="top"
                )
            widths[column_index - 1] = max(widths[column_index - 1], len(str(value)))
        for value_column, confidence_column in field_columns:
            value = row.get(value_column)
            confidence = row.get(confidence_column)
            if value not in {None, ""} and (
                value_column in row.get("_review_fields", ())
                or (
                    isinstance(confidence, (int, float))
                    and not isinstance(confidence, bool)
                    and confidence < low_confidence_threshold
                    and row.get(f"{value_column}_provenance_status") != "supplied"
                )
            ):
                sheet.cell(row=row_index, column=column_indexes[value_column]).fill = pale_yellow
                sheet.cell(row=row_index, column=column_indexes[confidence_column]).fill = pale_yellow
        if row.get("review_required"):
            cell = sheet.cell(row=row_index, column=column_indexes["review_required"])
            cell.fill = pale_yellow
            cell.font = Font(bold=True)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = min(max(width + 2, 12), 60)
    buffer = io.BytesIO()
    workbook.save(buffer)
    workbook.close()
    return buffer.getvalue()


def _xlsx_cell(value: Any) -> Any:
    if isinstance(value, str):
        # XML 1.0 cannot represent these control characters. Preserve other text,
        # replacing only invalid bytes with a visible replacement character.
        value = re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f\ud800-\udfff\ufffe\uffff]", "\ufffd", value)
    return _safe_cell(value)
