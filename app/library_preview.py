"""Bounded, read-only views of retained exports; never rerun research."""

from __future__ import annotations

import csv
import io
import math
import re
from collections.abc import Iterator
from typing import Any
from xml.etree.ElementTree import ParseError
from zipfile import BadZipFile, ZipFile

from openpyxl import load_workbook

from app.export.schema import RESEARCH_FIELDS, standard_columns
from app.library import SavedFileContent, SavedResultFile
from app.research.normalisation import normalise_display
from app.schemas import Contract

MAX_PREVIEW_COLUMNS = 512
MAX_PREVIEW_ROWS = 100_000
MAX_CELL_CHARS = 32_767


class LibraryPreviewError(ValueError):
    """The original artifact remains downloadable when a preview is unavailable."""


class SavedRecordsPage(Contract):
    file: SavedResultFile
    records: list[dict[str, Any]]
    total: int
    limit: int
    offset: int
    has_more: bool


def _key(value: Any) -> str:
    return re.sub(r"[^a-z0-9]+", "_", str(value or "").strip().lower()).strip("_")


def _research_columns(keys: list[str]) -> list[tuple[int, str]]:
    """Recognize the old export prefix before its appended input metadata.

    Older exports can contain both research ``status`` and administrative
    ``STATUS``. Only the known generated prefix (and optional provenance block)
    belongs to that format's research result; later input columns are not evidence.
    """
    fields = [field.value for field in RESEARCH_FIELDS]
    legacy_prefix = [
        *fields,
        *(f"{field}_confidence" for field in fields),
        "profile_confidence",
        "coverage",
        "legacy_coverage",
        "review_required",
        "status",
        "error_code",
        "source_urls",
    ]
    legacy_provenance = [
        *(f"{field}_source_urls" for field in fields),
        *(
            f"{field}_{detail}"
            for field in fields
            for detail in ("supplied_values", "input_confidence", "research_confidence", "provenance_status")
        ),
        "seed_input_metadata",
    ]
    legacy = keys[: len(legacy_prefix)] == legacy_prefix
    if legacy:
        stop = len(legacy_prefix)
        if keys[stop : stop + len(legacy_provenance)] == legacy_provenance:
            stop += len(legacy_provenance)
        keys = keys[:stop]
    known = set(standard_columns())
    selected = [(index, key) for index, key in enumerate(keys) if key in known]
    if len({key for _, key in selected}) != len(selected):
        raise LibraryPreviewError("Saved file has duplicate research columns")
    if legacy:
        # The old confidence column sometimes represented input certainty. Use
        # the explicit research value when the artifact actually retained it.
        for field in fields:
            research_key = f"{field}_research_confidence"
            if research_key in keys:
                target = f"{field}_confidence"
                selected = [(index, key) for index, key in selected if key != target]
                selected.append((keys.index(research_key), target))
    return selected


def _cell(value: Any) -> str | int | float | bool:
    if value is None:
        return ""
    if isinstance(value, (str, int, float, bool)):
        if isinstance(value, float) and not math.isfinite(value):
            return ""
        if isinstance(value, str) and len(value) > MAX_CELL_CHARS:
            raise LibraryPreviewError("A saved cell is too large to preview")
        return value
    return str(value)


def _csv_rows(content: bytes) -> Iterator[list[Any]]:
    # Each row is consumed independently; no dataframe or full-file row list.
    with io.TextIOWrapper(io.BytesIO(content), encoding="utf-8-sig", newline="") as stream:
        yield from csv.reader(stream, strict=True)


def _xlsx_rows(content: bytes, max_uncompressed_bytes: int) -> Iterator[list[Any]]:
    with ZipFile(io.BytesIO(content)) as archive:
        entries = archive.infolist()
        if len(entries) > 2048 or sum(item.file_size for item in entries) > max_uncompressed_bytes:
            raise LibraryPreviewError("Saved workbook is too large to preview")
    workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=False, keep_links=False)
    try:
        if not workbook.worksheets:
            raise LibraryPreviewError("Saved workbook has no records worksheet")
        sheet = workbook.worksheets[0]
        if sheet.max_column and sheet.max_column > MAX_PREVIEW_COLUMNS:
            raise LibraryPreviewError("Saved workbook has too many columns to preview")
        # Formula strings are returned as inert text. No formulas, macros or links
        # are evaluated, and the original bytes are never rewritten.
        for row in sheet.iter_rows(values_only=True):
            yield list(row)
    finally:
        workbook.close()


def preview_saved_file(
    saved: SavedFileContent,
    *,
    limit: int = 50,
    offset: int = 0,
    query: str = "",
    max_file_bytes: int = 10_000_000,
    max_uncompressed_bytes: int = 30_000_000,
) -> SavedRecordsPage:
    if not 1 <= limit <= 100 or offset < 0:
        raise LibraryPreviewError("Invalid record page")
    if len(saved.content) > max_file_bytes:
        raise LibraryPreviewError("Saved file is too large to preview")
    rows = (
        _csv_rows(saved.content)
        if saved.metadata.format == "csv"
        else _xlsx_rows(saved.content, max_uncompressed_bytes)
    )
    records: list[dict[str, Any]] = []
    matched = 0
    search = query.strip().casefold()
    try:
        header = next(rows, None)
        if not header or len(header) > MAX_PREVIEW_COLUMNS:
            raise LibraryPreviewError("Saved file has no usable research header")
        keys = [_key(value) for value in header]
        core = {field.value for field in RESEARCH_FIELDS}
        if not core.intersection(keys):
            raise LibraryPreviewError("Saved file has no recognized research columns")
        selected = _research_columns(keys)
        for index, cells in enumerate(rows):
            if index >= MAX_PREVIEW_ROWS:
                raise LibraryPreviewError("Saved file has too many rows to preview")
            if len(cells) > MAX_PREVIEW_COLUMNS:
                raise LibraryPreviewError("Saved file has too many columns to preview")
            if not any(value not in (None, "") for value in cells):
                continue
            record = {key: "" for key in standard_columns()}
            for position, key in selected:
                record[key] = _cell(cells[position]) if position < len(cells) else ""
            for field in RESEARCH_FIELDS:
                value = record.get(field.value)
                if isinstance(value, str) and value and not value.startswith(("=", "+", "@")):
                    record[field.value] = normalise_display(field, value)
            if search and not any(
                search in str(record.get(field.value, "")).casefold() for field in RESEARCH_FIELDS
            ):
                continue
            if offset <= matched < offset + limit:
                records.append(record)
            matched += 1
    except LibraryPreviewError:
        raise
    except (UnicodeError, csv.Error, BadZipFile, ValueError, KeyError, OSError, TypeError, ParseError) as exc:
        raise LibraryPreviewError("Saved file cannot be previewed; download the original artifact") from exc
    finally:
        rows.close()
    return SavedRecordsPage(
        file=saved.metadata,
        records=records,
        total=matched,
        limit=limit,
        offset=offset,
        has_more=offset + len(records) < matched,
    )
