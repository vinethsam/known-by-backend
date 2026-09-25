"""Offline end-to-end contracts through real providers, persistence and worker code."""

import asyncio
import csv
import io
import json
from uuid import uuid4

import httpx
import pytest
from alembic.config import Config
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook

from alembic import command
from app.api.routes import create_app
from app.config import Settings
from app.db import Store
from app.providers.openrouter import OpenRouterClient
from app.providers.search import OpenRouterSearchProvider
from app.research.orchestrator import ResearchOrchestrator
from app.retrieval.service import RetrievalService, RetrievedPage, StaticFetcher
from app.schemas import PersonSeed, ProfileField, SourceCandidate
from app.worker import process_lease, run_worker


def configured(tmp_path, **overrides):
    values = dict(
        APP_ENV="test",
        DATABASE_URL=f"sqlite:///{tmp_path / 'pipeline.db'}",
        OPENROUTER_API_KEY="fake-model-key",
        OPENROUTER_SOURCE_MODEL="test/source",
        OPENROUTER_EXTRACTION_MODEL="test/extraction",
        STATIC_FETCH_WORKER_URL="https://worker.example.net/fetch",
        STATIC_FETCH_WORKER_SECRET="fake-worker-key",
        API_ACCESS_TOKEN="test-token",
        PLAYWRIGHT_ENABLED=False,
        MAX_SEARCH_QUERIES_PER_PERSON=3,
        MAX_SOURCES_PER_PERSON=3,
        SOURCES_PER_ROUND=1,
        MAX_TOKENS_PER_PERSON=200000,
        WORKER_HEARTBEAT_SECONDS=0.05,
        WORKER_POLL_SECONDS=0.02,
    )
    values.update(overrides)
    return Settings(_env_file=None, **values)


def migrated_store(settings):
    config = Config("alembic.ini")
    config.set_main_option("sqlalchemy.url", settings.DATABASE_URL)
    command.upgrade(config, "head")
    return Store(settings)


def fake_services(
    settings,
    store,
    monkeypatch,
    *,
    malformed=False,
    search_urls=None,
    prose_only=False,
    search_status=200,
    always_new_query=False,
    malformed_citation=False,
    advisor_status=200,
    advisor_response=None,
    advisor_relevance="likely",
):
    calls = {
        "models": [],
        "queries": [],
        "fetched": [],
        "validated": [],
        "plans": 0,
        "search_limits": [],
        "advisor_requests": 0,
    }
    text = (
        "Jane Doe works for Example Foundation in Ghana as Programme Director. "
        "Jane Doe earned BSc in Economics at Example University."
    )
    urls = (
        search_urls
        if search_urls is not None
        else ["https://employer.example.org/jane", "https://university.example.edu/jane"]
    )

    async def safe(url, settings=None):
        return url

    monkeypatch.setattr("app.retrieval.service.validate_public_url", safe)

    def static(request):
        target = json.loads(request.content)["url"]
        calls["fetched"].append(target)
        assert request.headers["X-Worker-Secret"] == "fake-worker-key"
        return httpx.Response(
            200,
            json={
                "requested_url": target,
                "final_url": target,
                "status": 200,
                "content_type": "text/html",
                "html": f"<title>Jane Doe</title><h1>Jane Doe</h1><p>{text}</p><p>{target}</p>",
            },
        )

    def model(request):
        body = json.loads(request.content)
        calls["models"].append(body["model"])
        if body.get("tools"):
            calls["queries"].append(body["messages"][-1]["content"])
            calls["search_limits"].append(body["tools"][0]["parameters"]["max_results"])
            if malformed:
                return httpx.Response(200, content=b"{invalid")
            if search_status != 200:
                return httpx.Response(search_status, json={"error": {"message": "fixture failure"}})
            annotations = [
                {
                    "type": "url_citation",
                    "url_citation": {"url": url, "title": "Jane Doe", "content": text},
                }
                for url in urls
            ]
            if malformed_citation:
                annotations = [
                    {"type": "url_citation", "url_citation": {"title": "Missing URL"}},
                    {"type": "url_citation", "url_citation": "not-an-object"},
                    {"type": "other", "url_citation": {"url": "https://invented.example/fake"}},
                ]
            return httpx.Response(
                200,
                json={
                    "choices": [
                        {
                            "message": {
                                "content": "Model prose https://invented.example.org/fake",
                                "annotations": [] if prose_only else annotations,
                            }
                        }
                    ],
                    "usage": {
                        "prompt_tokens": 100,
                        "completion_tokens": 50,
                        "cost": 0.0002,
                        "server_tool_use": {"web_search_requests": 1},
                    },
                },
            )
        data = json.loads(body["messages"][1]["content"])["payload"]
        calls["advisor_requests"] += int(data.get("task") in {"plan_search_queries", "validate_candidates"})
        if data.get("task") == "plan_search_queries":
            calls["plans"] += 1
            answer = {
                "queries": [f'"Jane Doe" research {calls["plans"]}']
                if always_new_query
                else ['"Jane Doe" education', "education jane DOE", '"Jane Doe" Example Foundation']
            }
        elif data.get("task") == "validate_candidates":
            calls["validated"].extend(c["url"] for c in data["candidates"])
            if advisor_status != 200:
                return httpx.Response(
                    advisor_status,
                    json={"error": {"message": "provider rejected request fixture"}},
                )
            if advisor_response is not None:
                content = advisor_response
                return httpx.Response(
                    200,
                    json={
                        "choices": [{"message": {"content": content}}],
                        "usage": {"prompt_tokens": 100, "completion_tokens": 0, "cost": 0.0002},
                    },
                )
            answer = {
                "decisions": [
                    {
                        "candidate_id": c["candidate_id"],
                        "source_type": "employer",
                        "relevance": advisor_relevance,
                    }
                    for c in data["candidates"]
                ]
            }
        else:
            mapping = {
                "full_name": "Jane Doe",
                "organisation": "Example Foundation",
                "job_title": "Programme Director",
                "university_name": "Example University",
                "degree_type": "BSc",
                "subject": "Economics",
            }
            answer = {
                "claims": [
                    {
                        "field": f,
                        "value": v,
                        "evidence": text,
                        "subject_name": "Jane Doe",
                        "fact_group": "education"
                        if f in {"university_name", "degree_type", "subject"}
                        else "employment",
                    }
                    for f, v in mapping.items()
                ]
            }
        content = "{invalid" if malformed else json.dumps(answer)
        return httpx.Response(
            200,
            json={
                "choices": [{"message": {"content": content}}],
                "usage": {"prompt_tokens": 100, "completion_tokens": 50, "cost": 0.0002},
            },
        )

    model_client = httpx.AsyncClient(transport=httpx.MockTransport(model))
    fetch_client = httpx.AsyncClient(transport=httpx.MockTransport(static))
    shared_model = OpenRouterClient(settings, model_client, sleep=lambda _: asyncio.sleep(0))
    providers = (
        OpenRouterSearchProvider(settings, shared_model),
        shared_model,
        RetrievalService(settings, StaticFetcher(settings, fetch_client), cache=store),
    )
    return ResearchOrchestrator(settings, *providers), calls


@pytest.mark.asyncio
async def test_pipeline_round_trip_with_real_mocked_provider_clients(tmp_path, monkeypatch):
    settings = configured(tmp_path)
    store = migrated_store(settings)
    seed = PersonSeed(full_name="Jane Doe", organisation="Example Foundation", country="Ghana")
    job = store.create_job([seed])
    lease = store.claim_task("worker-1")
    pipeline, calls = fake_services(settings, store, monkeypatch)
    await process_lease(store, pipeline, lease, settings)
    finished = store.get_results(job.job_id)
    person = finished.people[0]
    assert person.status in {"completed", "review_required"}, person.error_code
    assert person.result.profile.fields[ProfileField.organisation].value == "Example Foundation"
    assert person.result.profile.fields[ProfileField.degree_type].value == "Bachelor's"
    assert person.result.profile.coverage == 100
    assert len(person.result.sources) == 2
    assert all(c.source_id in {s.source_id for s in person.result.sources} for c in person.result.claims)
    assert {"test/source", "test/extraction"} == set(calls["models"])
    assert len(calls["fetched"]) == len(set(calls["fetched"])) == 2
    assert sorted(calls["validated"]) == sorted(calls["fetched"])
    assert calls["advisor_requests"] >= 1
    assert person.result.profile.metrics.queries_performed <= 3
    assert person.result.profile.metrics.llm_calls == len(calls["models"])
    assert person.result.profile.metrics.tokens_used == len(calls["models"]) * 150
    assert len(person.result.usage) == len(calls["models"])


@pytest.mark.asyncio
async def test_linkedin_search_result_never_reaches_advisor_retrieval_or_evidence(tmp_path, monkeypatch):
    settings = configured(tmp_path, SOURCES_PER_ROUND=2)
    store = migrated_store(settings)
    job = store.create_job([PersonSeed(full_name="Jane Doe")])
    allowed = "https://official.example.org/jane"
    pipeline, calls = fake_services(
        settings,
        store,
        monkeypatch,
        search_urls=["https://www.linkedin.com/in/jane-doe", allowed],
    )

    await process_lease(store, pipeline, store.claim_task("worker"), settings)

    result = store.get_results(job.job_id).people[0].result
    assert calls["validated"] == calls["fetched"] == [allowed]
    assert [source.requested_url for source in result.sources] == [allowed]
    assert all("linkedin" not in source.final_url for source in result.sources)
    assert all(claim.source_id == result.sources[0].source_id for claim in result.claims)


@pytest.mark.asyncio
@pytest.mark.parametrize(
    "fixture_options",
    [
        {"prose_only": True},
        {"malformed_citation": True},
    ],
    ids=["prose-url-without-citations", "malformed-citation"],
)
async def test_no_usable_search_citations_finish_as_review_required(tmp_path, monkeypatch, fixture_options):
    settings = configured(
        tmp_path,
        MAX_SEARCH_QUERIES_PER_PERSON=3,
        MAX_SOURCE_MODEL_TOOL_CALLS=3,
    )
    store = migrated_store(settings)
    job = store.create_job([PersonSeed(full_name="Jane Doe")])
    pipeline, calls = fake_services(
        settings,
        store,
        monkeypatch,
        **fixture_options,
    )

    await process_lease(store, pipeline, store.claim_task("worker"), settings)

    person = store.get_results(job.job_id).people[0]
    assert person.status == "review_required"
    assert person.error_code is None
    assert person.result is not None
    assert person.result.profile.coverage == 0
    assert person.result.profile.research_status == "completed"
    assert person.result.profile.metrics.error_codes == []
    assert person.result.profile.metrics.stop_reason == "NO_SEARCH_CITATIONS"
    assert not calls["validated"] and not calls["fetched"]
    assert calls["plans"] == calls["advisor_requests"] == 0
    assert len(person.result.usage) == 1
    assert person.result.usage[0].success


@pytest.mark.asyncio
@pytest.mark.parametrize(
    ("fixture_options", "expected_code"),
    [
        ({"advisor_status": 400}, "SOURCE_ADVISOR_PROVIDER_ERROR"),
        ({"advisor_response": "{invalid"}, "SOURCE_ADVISOR_VALIDATION_ERROR"),
        (
            {"advisor_response": json.dumps({"decisions": "not-a-list"})},
            "SOURCE_ADVISOR_VALIDATION_ERROR",
        ),
    ],
    ids=["provider-http-400", "malformed-json", "schema-mismatch"],
)
async def test_source_advisor_failures_preserve_specific_terminal_code(
    tmp_path, monkeypatch, fixture_options, expected_code
):
    settings = configured(
        tmp_path,
        OPENROUTER_MAX_RETRIES=0,
        MAX_SEARCH_QUERIES_PER_PERSON=1,
    )
    store = migrated_store(settings)
    job = store.create_job([PersonSeed(full_name="Jane Doe")])
    pipeline, calls = fake_services(settings, store, monkeypatch, **fixture_options)
    attempt_logs, stage_logs = [], []
    monkeypatch.setattr(
        "app.research.budget.logger.warning",
        lambda event, *, extra: attempt_logs.append((event, extra)),
    )
    monkeypatch.setattr(
        "app.research.orchestrator.logger.error",
        lambda event, *, extra: stage_logs.append((event, extra)),
    )

    await process_lease(store, pipeline, store.claim_task("worker"), settings)

    person = store.get_results(job.job_id).people[0]
    assert person.status == "failed"
    assert person.error_code == expected_code
    assert person.result is not None
    assert expected_code in person.result.profile.metrics.error_codes
    assert calls["advisor_requests"] == 1
    assert not calls["fetched"]
    attempt = next(extra for event, extra in attempt_logs if event == "model_attempt")
    stage = next(extra for event, extra in stage_logs if event == "source_advisor_failed")
    assert attempt["operation"] == stage["operation"] == "validate_candidates"
    assert attempt["request_sent"] and attempt["response_received"]
    assert attempt["response_body_received"]
    assert attempt["error_code"] in {"OPENROUTER_HTTP_ERROR", "OPENROUTER_SCHEMA_ERROR"}
    assert stage["error_code"] == expected_code
    assert stage["candidate_count"] > 0 and stage["selected_source_count"] == 0
    assert "provider rejected request fixture" not in json.dumps([attempt, stage])


@pytest.mark.asyncio
async def test_all_unrelated_candidates_finish_without_retrieval(tmp_path, monkeypatch):
    settings = configured(
        tmp_path,
        MAX_SEARCH_QUERIES_PER_PERSON=3,
        MAX_SOURCE_MODEL_TOOL_CALLS=3,
    )
    store = migrated_store(settings)
    job = store.create_job([PersonSeed(full_name="Jane Doe")])
    pipeline, calls = fake_services(
        settings,
        store,
        monkeypatch,
        advisor_relevance="unrelated",
    )

    await process_lease(store, pipeline, store.claim_task("worker"), settings)

    person = store.get_results(job.job_id).people[0]
    assert person.status == "review_required"
    assert person.error_code is None
    assert person.result is not None
    assert person.result.profile.research_status == "completed"
    assert person.result.profile.metrics.error_codes == []
    assert person.result.profile.metrics.stop_reason == "NO_SELECTED_SOURCES"
    assert calls["validated"] and not calls["fetched"]


@pytest.mark.asyncio
async def test_candidates_empty_after_filtering_finish_without_advisor(tmp_path, monkeypatch):
    settings = configured(tmp_path, MAX_SEARCH_QUERIES_PER_PERSON=3)
    store = migrated_store(settings)
    job = store.create_job([PersonSeed(full_name="Jane Doe")])
    pipeline, calls = fake_services(settings, store, monkeypatch)
    monkeypatch.setattr(
        pipeline.search,
        "_parse_results",
        lambda _data, _limit: [SourceCandidate(url="not-a-valid-public-url")],
    )

    await process_lease(store, pipeline, store.claim_task("worker"), settings)

    person = store.get_results(job.job_id).people[0]
    assert person.status == "review_required"
    assert person.error_code is None
    assert person.result is not None
    assert person.result.profile.metrics.stop_reason == "NO_ELIGIBLE_CANDIDATES"
    assert calls["advisor_requests"] == 0
    assert not calls["validated"] and not calls["fetched"]


@pytest.mark.asyncio
async def test_budget_exhaustion_before_advisor_makes_no_advisor_request(tmp_path, monkeypatch):
    settings = configured(
        tmp_path,
        MAX_LLM_CALLS_PER_PERSON=1,
        MAX_SEARCH_QUERIES_PER_PERSON=1,
    )
    store = migrated_store(settings)
    pipeline, calls = fake_services(settings, store, monkeypatch)

    result = await pipeline.research(str(uuid4()), str(uuid4()), PersonSeed(full_name="Jane Doe"))

    assert result.profile.metrics.stop_reason == "MAX_LLM_CALLS"
    assert calls["advisor_requests"] == 0
    assert not calls["validated"] and not calls["fetched"]
    assert len(calls["models"]) == len(result.usage) == 1
    assert result.usage[0].success


@pytest.mark.asyncio
async def test_batch_worker_reuses_retrieval_and_stops(tmp_path, monkeypatch):
    settings = configured(tmp_path, MAX_CONCURRENT_PEOPLE=2)
    store = migrated_store(settings)
    seed = PersonSeed(full_name="Jane Doe", organisation="Example Foundation", country="Ghana")
    job = store.create_job([seed, seed])
    pipeline, calls = fake_services(settings, store, monkeypatch)
    stop = asyncio.Event()
    worker = asyncio.create_task(run_worker(settings, store, pipeline, stop))
    try:
        async with asyncio.timeout(10):
            while (await asyncio.to_thread(store.get_job, job.job_id)).status in {"queued", "running"}:
                await asyncio.sleep(0.02)
    finally:
        stop.set()
        await worker
    assert store.get_job(job.job_id).status == "completed"
    assert len(calls["fetched"]) == 2
    assert all(p.result.profile.coverage == 100 for p in store.get_results(job.job_id).people)


@pytest.mark.asyncio
async def test_paid_retries_obey_person_call_budget(tmp_path, monkeypatch):
    settings = configured(tmp_path, MAX_LLM_CALLS_PER_PERSON=1)
    store = migrated_store(settings)
    pipeline, calls = fake_services(settings, store, monkeypatch, malformed=True)
    result = await pipeline.research(str(uuid4()), str(uuid4()), PersonSeed(full_name="Jane Doe"))
    assert len(calls["models"]) == 1
    assert result.profile.metrics.stop_reason == "MAX_LLM_CALLS"
    assert result.profile.coverage == 0
    assert len(result.usage) == 1 and not result.usage[0].success


@pytest.mark.asyncio
async def test_search_retries_reserve_tool_calls_before_paid_requests(tmp_path, monkeypatch):
    settings = configured(tmp_path, MAX_SOURCE_MODEL_TOOL_CALLS=1, OPENROUTER_MAX_RETRIES=2)
    store = migrated_store(settings)
    pipeline, calls = fake_services(settings, store, monkeypatch, search_status=429)
    result = await pipeline.research(str(uuid4()), str(uuid4()), PersonSeed(full_name="Jane Doe"))
    assert len(calls["queries"]) == 1
    assert result.profile.metrics.web_search_calls_reserved == 1
    assert result.profile.metrics.stop_reason == "MAX_SEARCH_TOOL_CALLS"
    assert len(result.usage) == 1 and not result.usage[0].success


@pytest.mark.asyncio
async def test_search_result_reservations_cap_request_and_reject_prose_urls(tmp_path, monkeypatch):
    settings = configured(tmp_path, MAX_TOTAL_SEARCH_RESULTS_PER_PERSON=3)
    store = migrated_store(settings)
    pipeline, calls = fake_services(settings, store, monkeypatch, prose_only=True)
    result = await pipeline.research(str(uuid4()), str(uuid4()), PersonSeed(full_name="Jane Doe"))
    assert calls["search_limits"] == [3]
    assert not calls["fetched"] and not calls["validated"]
    assert result.profile.metrics.search_results_reserved == 3
    assert result.profile.metrics.stop_reason == "NO_SEARCH_CITATIONS"
    assert len(calls["queries"]) == 1
    assert calls["plans"] == calls["advisor_requests"] == 0
    assert result.profile.coverage == 0


@pytest.mark.asyncio
async def test_zero_citation_search_stops_without_advisor_and_usage_persists(tmp_path, monkeypatch):
    settings = configured(tmp_path, MAX_SEARCH_QUERIES_PER_PERSON=2, MAX_TOTAL_SEARCH_RESULTS_PER_PERSON=100)
    store = migrated_store(settings)
    pipeline, calls = fake_services(settings, store, monkeypatch, prose_only=True, always_new_query=True)
    job = store.create_job([PersonSeed(full_name="Jane Doe")])
    await process_lease(store, pipeline, store.claim_task("worker"), settings)
    result = store.get_results(job.job_id).people[0].result
    assert len(calls["queries"]) == result.profile.metrics.queries_performed == 1
    assert calls["plans"] == calls["advisor_requests"] == 0
    assert result.profile.metrics.stop_reason == "NO_SEARCH_CITATIONS"
    assert sum(u.web_search_requests or 0 for u in result.usage) == 1
    assert result.profile.metrics.web_search_requests == 1


@pytest.mark.asyncio
async def test_discovered_backlog_is_used_before_searching_and_classification_is_reused(
    tmp_path, monkeypatch
):
    settings = configured(tmp_path, MAX_SOURCE_MODEL_TOOL_CALLS=1, SOURCES_PER_ROUND=1)
    store = migrated_store(settings)
    urls = ["https://employer.example.org/jane", "https://university.example.edu/jane"]
    pipeline, calls = fake_services(
        settings, store, monkeypatch, search_urls=urls + [urls[0] + "?utm_source=copy"]
    )
    result = await pipeline.research(
        str(uuid4()),
        str(uuid4()),
        PersonSeed(full_name="Jane Doe", organisation="Example Foundation", country="Ghana"),
    )
    assert len(calls["queries"]) == 1
    assert calls["plans"] == 0
    assert sorted(calls["fetched"]) == sorted(urls)
    assert sorted(calls["validated"]) == sorted(urls)
    assert result.profile.coverage == 100


@pytest.mark.asyncio
async def test_discovery_and_extraction_can_use_same_model(tmp_path, monkeypatch):
    settings = configured(
        tmp_path, OPENROUTER_SOURCE_MODEL="test/shared", OPENROUTER_EXTRACTION_MODEL="test/shared"
    )
    store = migrated_store(settings)
    pipeline, calls = fake_services(settings, store, monkeypatch)
    result = await pipeline.research(
        str(uuid4()),
        str(uuid4()),
        PersonSeed(full_name="Jane Doe", organisation="Example Foundation", country="Ghana"),
    )
    assert set(calls["models"]) == {"test/shared"}
    assert result.profile.coverage == 100


@pytest.mark.asyncio
async def test_search_is_fenced_by_token_budget_before_http(tmp_path, monkeypatch):
    settings = configured(tmp_path, MAX_TOKENS_PER_PERSON=1000)
    store = migrated_store(settings)
    pipeline, calls = fake_services(settings, store, monkeypatch)
    result = await pipeline.research(str(uuid4()), str(uuid4()), PersonSeed(full_name="Jane Doe"))
    assert not calls["models"]
    assert result.profile.metrics.stop_reason == "TOKEN_BUDGET"


@pytest.mark.asyncio
async def test_search_citation_dns_still_passes_production_retrieval_guard(tmp_path, monkeypatch):
    import socket

    from app.retrieval.urls import validate_public_url

    settings = configured(tmp_path, MAX_SOURCE_MODEL_TOOL_CALLS=1)
    store = migrated_store(settings)
    pipeline, calls = fake_services(
        settings, store, monkeypatch, search_urls=["https://employer.example.org/jane"]
    )
    monkeypatch.setattr("app.retrieval.service.validate_public_url", validate_public_url)
    monkeypatch.setattr(
        socket,
        "getaddrinfo",
        lambda host, port, **kwargs: [(socket.AF_INET, socket.SOCK_STREAM, 6, "", ("127.0.0.1", port))],
    )
    job = store.create_job([PersonSeed(full_name="Jane Doe")])
    await process_lease(store, pipeline, store.claim_task("worker"), settings)
    person = store.get_results(job.job_id).people[0]
    result = person.result
    assert person.status == "failed"
    assert person.error_code == "RETRIEVAL_FAILED"
    assert result is not None
    assert calls["validated"] and not calls["fetched"]
    assert "RETRIEVAL_FAILED" in result.profile.metrics.error_codes


class FakeFetcher:
    async def fetch(self, url):
        return RetrievedPage(
            requested_url=url,
            final_url=url,
            status=200,
            content_type="text/html",
            html="<h1>Preserved</h1><p>A <a href='/profile'>profile</a>.</p>",
        )

    async def close(self):
        pass


def api_client(tmp_path, **overrides):
    settings = configured(tmp_path, **overrides)
    store = migrated_store(settings)
    return TestClient(create_app(settings, store, FakeFetcher())), store


AUTH = {"Authorization": "Bearer test-token"}


def test_api_health_ready_auth_and_immediate_job(tmp_path):
    client, store = api_client(tmp_path)
    with client:
        assert client.get("/health").json() == {"status": "ok"}
        assert client.get("/ready").status_code == 503
        store.heartbeat("worker")
        assert client.get("/ready").status_code == 200
        assert client.post("/v1/research/person", json={"full_name": "Jane Doe"}).status_code == 401
        response = client.post("/v1/research/person", headers=AUTH, json={"full_name": "Jane Doe"})
        assert response.status_code == 202
        job_id = response.json()["job_id"]
        status = client.get(f"/v1/jobs/{job_id}", headers=AUTH).json()
        assert status["status"] == "queued" and status["counts"]["queued"] == 1
        results = client.get(f"/v1/jobs/{job_id}/results", headers=AUTH).json()
        assert results["people"][0]["result"] is None
        assert client.get(f"/v1/jobs/{uuid4()}", headers=AUTH).status_code == 404
        assert client.post(f"/v1/jobs/{job_id}/cancel", headers=AUTH).json()["status"] == "cancelled"


@pytest.mark.parametrize("kind", ["csv", "xlsx"])
def test_api_batch_and_export_contracts(tmp_path, kind):
    client, store = api_client(tmp_path)
    content = b"Person,Department\nJane Doe,Research\nJohn Smith,Engineering\n"
    if kind == "xlsx":
        workbook = Workbook()
        for row in [["Person", "Department"], ["Jane Doe", "Research"], ["John Smith", "Engineering"]]:
            workbook.active.append(row)
        stream = io.BytesIO()
        workbook.save(stream)
        content = stream.getvalue()
    with client:
        response = client.post(
            "/v1/research/batch",
            headers=AUTH,
            files={"file": (f"people.{kind}", content)},
        )
        assert response.status_code == 202, response.text
        job_id = response.json()["job_id"]
        assert response.json()["total_people"] == 2
        result = client.get(f"/v1/jobs/{job_id}/results", headers=AUTH).json()
        assert result["people"][0]["original_row"]["Department"] == "Research"
        for output in ("csv", "xlsx"):
            export = client.get(f"/v1/jobs/{job_id}/export?format={output}", headers=AUTH)
            assert export.status_code == 200
            if output == "xlsx":
                sheet = load_workbook(io.BytesIO(export.content)).active
                assert sheet.cell(2, 1).value == "Jane Doe"
            else:
                assert "Department" in export.text and "Jane Doe" in export.text
        rich_export = client.get(f"/v1/jobs/{job_id}/export?format=csv&provenance=field", headers=AUTH)
        assert rich_export.status_code == 200
        rich_headers = next(csv.reader(io.StringIO(rich_export.content.decode("utf-8-sig"))))
        assert "full_name_source_urls" in rich_headers
        assert rich_headers.index("full_name_source_urls") == rich_headers.index("full_name_confidence") + 1
        assert client.get(f"/v1/jobs/{job_id}/export?provenance=claims", headers=AUTH).status_code == 422


def test_api_rejects_unsafe_urls_and_missing_configuration(tmp_path):
    client, _ = api_client(tmp_path)
    with client:
        response = client.post(
            "/v1/research/person",
            headers=AUTH,
            json={"full_name": "Jane Doe", "preferred_urls": ["http://127.0.0.1/secret"]},
        )
        assert response.status_code == 422
        response = client.post(
            "/v1/research/person", headers=AUTH, json={"full_name": "J", "secret": "do-not-echo"}
        )
        assert response.status_code == 422 and "do-not-echo" not in response.text
    settings = configured(tmp_path, OPENROUTER_API_KEY=None)
    with TestClient(create_app(settings, Store(settings), FakeFetcher())) as unavailable:
        assert (
            unavailable.post("/v1/research/person", headers=AUTH, json={"full_name": "Jane Doe"}).status_code
            == 503
        )


def test_legacy_process_contract_and_upload_limit(tmp_path):
    client, _ = api_client(tmp_path, MAX_UPLOAD_BYTES=1024)
    with client:
        result = client.post("/process", headers=AUTH, json={"url": "https://example.org"})
        assert result.status_code == 200
        assert "# Preserved" in result.json()["markdown"]
        assert "https://example.org/profile" in result.json()["markdown"]
        oversized = client.post("/v1/research/batch", headers=AUTH, files={"file": ("x.csv", b"x" * 2000)})
        assert oversized.status_code == 413
