"""Offline text-boundary regressions; fixtures contain no real input records."""

import csv
import io
import json

import pytest
from openpyxl import Workbook, load_workbook

from app.config import Settings
from app.export.batch import BatchParseError, parse_batch
from app.export.formats import export_results
from app.input_validation import sanitize_input_text
from app.research.discovery import build_context_queries
from app.research.identity import research_seed_payload
from app.research.reconciliation import reconcile
from app.schemas import JobResults, PersonResultView, PersonSeed, ResearchResult
from app.seed_context import seed_context_values


def csv_bytes(rows):
    output = io.StringIO(newline="")
    csv.writer(output).writerows(rows)
    return output.getvalue().encode("utf-8")


@pytest.mark.parametrize(
    "raw,expected",
    [
        ("Jo\x00sé Example", "José Example"),
        ("Ada\vExample", "Ada Example"),
        ("Ada\fExample", "Ada Example"),
        ("Ada\x85Example", "Ada Example"),
        ("Ada\x7f Example\x81", "Ada Example"),
        ("\u00c3\u0081lvaro Example", "Álvaro Example"),
        ("José 王 සිංහ\u200d نادیا\u200c", "José 王 සිංහ\u200d نادیا\u200c"),
        ("  two  spaces\u00a0remain  ", "  two  spaces\u00a0remain  "),
        ("Ａda Ⅳ Example", "Ａda Ⅳ Example"),
    ],
)
def test_central_sanitation_repairs_before_removing_controls_and_preserves_unicode(raw, expected):
    assert sanitize_input_text(raw) == expected
    assert sanitize_input_text(expected) == expected


def test_multiline_sanitation_keeps_meaningful_layout():
    raw = "one\tcolumn\nsecond\r\nthird\rfourth\vnext\fpage\x85last\x00"
    assert sanitize_input_text(raw, multiline=True) == (
        "one\tcolumn\nsecond\r\nthird\rfourth\nnext\npage\nlast"
    )
    assert sanitize_input_text(raw) == "one column second  third fourth next page last"


@pytest.mark.parametrize("control", ["\x00", "\v", "\f", "\x7f", "\x81"])
def test_csv_cells_preserve_raw_input_and_use_only_clean_processing_text(control):
    original = {
        "Full Name": f"Áda{control} Example",
        "University": f"Lake{control} University",
        "Subject": "Physics\nAstronomy",
        "Notes": "original\tline\nnext\r\nlast" + control,
        "Email": "private@example.test",
    }
    batch = parse_batch(
        csv_bytes([list(original), list(original.values())]),
        "people.csv",
        None,
        Settings(_env_file=None, APP_ENV="test"),
    )
    seed = batch.seeds[0]
    assert seed.full_name == sanitize_input_text(original["Full Name"])
    assert seed.university_name == sanitize_input_text(original["University"])
    assert seed.subject == "Physics Astronomy"
    assert seed.input_metadata == original
    assert batch.rows == [original]
    assert PersonSeed.model_validate_json(seed.model_dump_json()).input_metadata == original
    context = seed_context_values(seed)
    assert context["university_name"] == ["Lake University"]
    queries = " ".join(build_context_queries(seed))
    payload = json.dumps(research_seed_payload(seed), ensure_ascii=False)
    for text in [queries, payload]:
        assert control not in text
        assert "private@example.test" not in text
    assert "Lake University" in queries


@pytest.mark.parametrize("format", ["csv", "xlsx"])
def test_misencoded_accent_that_triggered_baseline_failure_is_repaired_without_losing_raw_name(format):
    original = {"ALUMNI_FULL_NAME": "\u00c3\u0081lvaro Example", "ALUMNI_FIRST_NAME": "\u00c3\u0081lvaro"}
    if format == "csv":
        content = csv_bytes([list(original), list(original.values())])
    else:
        workbook = Workbook()
        workbook.active.append(list(original))
        workbook.active.append(list(original.values()))
        output = io.BytesIO()
        workbook.save(output)
        workbook.close()
        content = output.getvalue()
    batch = parse_batch(content, f"people.{format}", None, Settings(_env_file=None, APP_ENV="test"))
    assert batch.seeds[0].full_name == "Álvaro Example"
    assert batch.rows == [original]
    assert batch.seeds[0].input_metadata == original


def test_raw_metadata_aliases_cannot_bypass_processing_sanitation():
    seed = PersonSeed(
        full_name="Ada Example",
        input_metadata={"University": "Lake\x00 University", "Previous Employer": "Old\vLabs"},
    )
    assert seed_context_values(seed) == {
        "university_name": ["Lake University"],
        "historical_employer": ["Old Labs"],
    }
    assert seed.input_metadata["University"] == "Lake\x00 University"


@pytest.mark.parametrize("format", ["csv", "xlsx"])
def test_provenance_round_trip_preserves_original_controls_without_creating_evidence(format):
    raw = "Lake\x00 University\vFaculty\fPhysics"
    seed = PersonSeed(full_name="Ada Example", university_name=raw)
    cfg = Settings(_env_file=None, APP_ENV="test")
    profile = reconcile("person", seed, [], [], cfg.SCORING)
    assert profile.coverage == 0
    results = JobResults(
        job_id="job",
        status="completed",
        people=[
            PersonResultView(
                person_id="person",
                row_index=2,
                original_row=seed.input_metadata,
                status="review_required",
                result=ResearchResult(profile=profile),
            )
        ],
    )
    output = export_results(results, list(seed.input_metadata), format, provenance="field")
    if format == "csv":
        row = next(csv.DictReader(io.StringIO(output.decode("utf-8-sig"))))
    else:
        workbook = load_workbook(io.BytesIO(output), read_only=True)
        rows = workbook.active.iter_rows(values_only=True)
        row = dict(zip(next(rows), next(rows), strict=True))
        workbook.close()
    assert "seed_input_metadata" not in row
    attributes = profile.seed_attributes
    supplied = next(item for item in attributes if item["field"] == "university_name")
    assert supplied["value"] == raw
    assert supplied["input_confidence"] == 100
    assert json.loads(row["university_name_supplied_values"]) == [raw]
    assert float(row["university_name_confidence"]) == 0
    assert row["university_name"] == "Lake University Faculty Physics"


def test_header_controls_are_sanitized_for_matching_without_losing_original_labels():
    batch = parse_batch(
        b"Full\x00 Name,University\nAda Example,Lake University\n",
        "people.csv",
        None,
        Settings(_env_file=None, APP_ENV="test"),
    )
    assert batch.seeds[0].full_name == "Ada Example"
    assert batch.columns == ["Full\x00 Name", "University"]
    with pytest.raises(BatchParseError, match="duplicate"):
        parse_batch(
            b"name,na\x00me\nAda Example,Ada Example\n",
            "people.csv",
            None,
            Settings(_env_file=None, APP_ENV="test"),
        )
