from __future__ import annotations

import csv
import io
import re
import zipfile

import pytest
from openpyxl import Workbook, load_workbook

from app.config import Settings
from app.export.batch import BatchParseError, normalize_header, parse_batch
from app.export.formats import export_results
from app.schemas import (
    FieldDecision,
    JobResults,
    PersonProfile,
    PersonResultView,
    PersonStatus,
    ProfileField,
    ProfileRecord,
    ResearchResult,
    SourceRecord,
)


def _settings(**overrides):
    data = {
        "APP_ENV": "test",
        "MAX_BATCH_ROWS": 2,
        "MAX_UPLOAD_BYTES": 20_000,
        "MAX_XLSX_UNCOMPRESSED_BYTES": 100_000,
        "MAX_INPUT_COLUMNS": 5,
    }
    data.update(overrides)
    return Settings(**data)


def test_parse_csv_preserves_rows_columns_and_row_indexes():
    content = b"name,organisation,note\nAda Lovelace,Analytical Engines,first\nGrace Hopper,Navy,second\n"
    batch = parse_batch(content, "people.csv", None, _settings())
    assert [seed.full_name for seed in batch.seeds] == ["Ada Lovelace", "Grace Hopper"]
    assert batch.columns == ["name", "organisation", "note"]
    assert batch.rows[0] == {
        "name": "Ada Lovelace",
        "organisation": "Analytical Engines",
        "note": "first",
    }
    assert batch.seeds[0].organisation == "Analytical Engines"
    assert batch.seeds[0].known_attributes == {}


@pytest.mark.parametrize(
    "header",
    [
        "Full Name",
        "fullname",
        "full_name",
        "FULL-NAME",
        "fullName",
        "full.name",
        "  Full---Name  ",
        "Ｆｕｌｌ　Ｎａｍｅ",
    ],
)
def test_header_normalization_handles_common_schema_styles(header):
    assert normalize_header(header) == "full name"


@pytest.mark.parametrize(
    "header",
    [
        "Full Name",
        "fullname",
        "alumni_full_name",
        "memberName",
        "candidate_name",
        "person",
        "official-name",
    ],
)
def test_flexible_person_name_aliases_create_seeds(header):
    batch = parse_batch(f"{header}\nAda Lovelace\n".encode(), "people.csv", None, _settings())

    assert [seed.full_name for seed in batch.seeds] == ["Ada Lovelace"]


@pytest.mark.parametrize(
    ("context_header", "value", "seed_field"),
    [
        ("company_name", "Analytical Engines", "organisation"),
        ("university_name", "Example University", "university_name"),
        ("government_office", "Ministry of Science", "organisation"),
        ("job_title", "Director", "job_title"),
        ("country_of_origin", "Kenya", "country"),
        ("currentLocation", "Nairobi", "location"),
        ("field_of_study", "Economics", "subject"),
        ("graduation_year", "2020", "program_year"),
    ],
)
def test_context_headers_map_to_the_correct_seed_field(context_header, value, seed_field):
    content = f"candidate_name,{context_header}\nJane Doe,{value}\n".encode()
    seed = parse_batch(content, "people.csv", None, _settings()).seeds[0]

    assert seed.full_name == "Jane Doe"
    assert getattr(seed, seed_field) == value
    if context_header in {"company_name", "university_name"}:
        wrong_field = "university_name" if seed_field == "organisation" else "organisation"
        assert getattr(seed, wrong_field) is None


def test_flexible_schema_uses_context_and_ignores_unknown_columns():
    content = b"name,company,country,favorite_colour,random_code\nJane Doe,Ministry X,Kenya,blue,ABC-123\n"
    batch = parse_batch(content, "people.csv", None, _settings())
    seed = batch.seeds[0]

    assert seed.full_name == "Jane Doe"
    assert seed.organisation == "Ministry X"
    assert seed.country == "Kenya"
    assert seed.known_attributes == {}
    assert batch.rows[0]["favorite_colour"] == "blue"
    assert batch.rows[0]["random_code"] == "ABC-123"


def test_more_specific_person_name_alias_beats_generic_name():
    batch = parse_batch(
        b"name,alumni_full_name\nAda,Adelaide Lovelace\n",
        "people.csv",
        None,
        _settings(),
    )

    assert batch.seeds[0].full_name == "Adelaide Lovelace"


def test_explicit_name_column_override_remains_backward_compatible():
    batch = parse_batch(
        b"display_label,company\nAda Lovelace,Analytical Engines\n",
        "people.csv",
        "display_label",
        _settings(),
    )

    assert batch.seeds[0].full_name == "Ada Lovelace"
    assert batch.seeds[0].organisation == "Analytical Engines"


def test_equal_person_name_candidates_report_the_headers():
    with pytest.raises(BatchParseError, match="ambiguous name columns") as error:
        parse_batch(
            b"candidate_name,member_name\nJane Doe,Jane Doe\n",
            "people.csv",
            None,
            _settings(),
        )

    assert "candidate_name" in str(error.value)
    assert "member_name" in str(error.value)


def test_missing_name_error_lists_observed_headers():
    with pytest.raises(BatchParseError, match="missing a name column") as error:
        parse_batch(
            b"company_name,random_code\nExample Ltd,123\n",
            "people.csv",
            None,
            _settings(),
        )

    assert "company_name" in str(error.value)
    assert "random_code" in str(error.value)


@pytest.mark.parametrize("value", ["123456", "https://example.org/profile", "person@example.org", ""])
def test_obvious_non_name_columns_are_not_accepted(value):
    content = f"candidate,company\n{value},Example Ltd\n".encode()

    with pytest.raises(BatchParseError, match="missing a name column"):
        parse_batch(content, "people.csv", None, _settings())


@pytest.mark.parametrize(
    ("content", "message"),
    [
        (b"name,name\nAda,Again\n", "duplicate"),
        (b"name,\nAda,x\n", "blank"),
        (b"favorite_colour\nblue\n", "missing"),
        (b"candidate_name,member_name\nAda Lovelace,Ada Lovelace\n", "ambiguous"),
        (b"name\n\n", "Blank row"),
    ],
)
def test_parse_csv_rejects_bad_headers_and_missing_people(content, message):
    with pytest.raises(BatchParseError, match=message):
        parse_batch(content, "people.csv", None, _settings())


def test_parse_rejects_limits_and_missing_name():
    with pytest.raises(BatchParseError, match="too many rows"):
        parse_batch(b"name\nAda Lovelace\nGrace Hopper\nAlan Turing\n", "people.csv", None, _settings())
    with pytest.raises(BatchParseError, match="Missing name"):
        parse_batch(b"name\nAda Lovelace\n \n", "people.csv", None, _settings(MAX_BATCH_ROWS=5))
    with pytest.raises(BatchParseError, match="too many columns"):
        parse_batch(b"name,a,b,c,d,e\nAda,1,2,3,4,5\n", "people.csv", None, _settings())
    with pytest.raises(BatchParseError, match="maximum size"):
        parse_batch(b"name\n" + (b"a" * 2000), "people.csv", None, _settings(MAX_UPLOAD_BYTES=1024))


def test_parse_xlsx_preserves_formula_text_and_rejects_zipbomb(tmp_path):
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["full name", "note"])
    sheet.append(["Ada Lovelace", "=1+1"])
    buffer = io.BytesIO()
    workbook.save(buffer)
    workbook.close()

    batch = parse_batch(buffer.getvalue(), "people.xlsx", None, _settings())
    assert batch.rows[0]["note"] == "=1+1"

    bomb = io.BytesIO()
    with zipfile.ZipFile(bomb, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("xl/worksheets/sheet1.xml", "x" * 200_000)
    with pytest.raises(BatchParseError, match="maximum allowed size"):
        parse_batch(bomb.getvalue(), "people.xlsx", None, _settings(MAX_XLSX_UNCOMPRESSED_BYTES=1024))


def _results() -> JobResults:
    person_id = "person-1"
    fields = {
        field: FieldDecision(value=f"{field.value} value", confidence=88, review_required=False)
        for field in ProfileField
    }
    fields[ProfileField.full_name] = FieldDecision(value="=Ada", confidence=90, review_required=False)
    profile = PersonProfile(
        person_id=person_id,
        input_name="Ada Lovelace",
        status=PersonStatus.completed,
        fields=fields,
        profile_confidence=92,
        coverage=100,
        review_required=False,
    )
    source = SourceRecord(
        person_id=person_id,
        requested_url="https://example.com",
        final_url="https://example.com",
        canonical_url="https://example.com",
        domain="example.com",
    )
    return JobResults(
        job_id="job-1",
        status="completed",
        people=[
            PersonResultView(
                person_id=person_id,
                row_index=2,
                original_row={"name": "+Ada Lovelace", "status": "original"},
                status=PersonStatus.completed,
                result=ResearchResult(profile=profile, sources=[source], claims=[], usage=[]),
            )
        ],
    )


def test_export_csv_is_formula_safe_and_collision_safe():
    data = export_results(_results(), ["name", "status"], "csv")
    rows = list(csv.DictReader(io.StringIO(data.decode("utf-8-sig"))))
    assert rows[0]["name"] == "'+Ada Lovelace"
    assert rows[0]["status"] == "original"
    assert rows[0]["status (2)"] == "completed"
    assert rows[0]["full_name"] == "'=Ada"
    assert rows[0]["source_urls"] == "https://example.com"


def test_export_xlsx_is_readable_and_formula_safe():
    data = export_results(_results(), ["name", "status"], "xlsx")
    workbook = load_workbook(io.BytesIO(data), data_only=False)
    sheet = workbook.active
    headers = [cell.value for cell in sheet[1]]
    row = [cell.value for cell in sheet[2]]
    workbook.close()
    assert headers[:3] == ["name", "status", "status (2)"]
    assert row[0] == "'+Ada Lovelace"
    assert row[headers.index("full_name")] == "'=Ada"


def test_default_export_is_unchanged_for_one_record_projection():
    results = _results()
    expected = export_results(results, ["name", "status"], "csv")
    profile = results.people[0].result.profile
    record = ProfileRecord(
        record_id="primary",
        fields=profile.fields,
        profile_confidence=profile.profile_confidence,
        coverage=profile.coverage,
        review_required=profile.review_required,
        status=profile.status,
    )
    profile.records = [record]

    assert export_results(results, ["name", "status"], "csv") == expected


@pytest.mark.parametrize("format", ["csv", "xlsx"])
def test_field_provenance_export_expands_records_without_cross_record_sources(format):
    results = _results()
    profile = results.people[0].result.profile

    def record(degree, university, source_url, profile_url, confidence):
        fields = {field: decision.model_copy(deep=True) for field, decision in profile.fields.items()}
        fields[ProfileField.degree_type] = FieldDecision(
            value=degree,
            confidence=confidence,
            sources=[source_url, source_url],
            review_required=False,
        )
        fields[ProfileField.university_name] = FieldDecision(
            value=university,
            confidence=confidence,
            sources=[source_url],
            review_required=False,
        )
        fields[ProfileField.profile_link] = FieldDecision(
            value=profile_url,
            confidence=confidence,
            sources=[profile_url],
            review_required=False,
        )
        return ProfileRecord(
            record_id=degree.casefold(),
            fields=fields,
            profile_confidence=confidence,
            coverage=100,
            review_required=False,
            status=PersonStatus.completed,
        )

    bachelor_url = "https://university-a.example/ada"
    master_url = "https://university-b.example/ada"
    profile.records = [
        record("Bachelor's", "University A", bachelor_url, bachelor_url, 91),
        record("Master's", "University B", master_url, master_url, 89),
    ]

    data = export_results(results, ["name", "status"], format, provenance="field")
    if format == "csv":
        rows = list(csv.DictReader(io.StringIO(data.decode("utf-8-sig"))))
        headers = list(rows[0])
    else:
        workbook = load_workbook(io.BytesIO(data), data_only=False)
        values = list(workbook.active.values)
        workbook.close()
        headers = list(values[0])
        rows = [dict(zip(headers, values[index], strict=True)) for index in range(1, len(values))]

    assert len(rows) == 2
    assert headers.index("degree_type_source_urls") == headers.index("degree_type_confidence") + 1
    assert [row["degree_type"] for row in rows] == ["Bachelor's", "Master's"]
    assert [row["name"] for row in rows] == ["'+Ada Lovelace", "'+Ada Lovelace"]
    assert rows[0]["degree_type_source_urls"] == bachelor_url
    assert rows[1]["degree_type_source_urls"] == master_url
    assert master_url not in rows[0]["university_name_source_urls"]
    assert bachelor_url not in rows[1]["university_name_source_urls"]
    assert [row["profile_link"] for row in rows] == [bachelor_url, master_url]


def test_field_provenance_export_keeps_failed_person_and_formula_safety():
    results = _results()
    results.people[0].result = None
    results.people[0].status = PersonStatus.failed
    failed = list(
        csv.DictReader(
            io.StringIO(export_results(results, ["name"], "csv", provenance="field").decode("utf-8-sig"))
        )
    )
    assert len(failed) == 1
    assert failed[0]["status"] == "failed"
    assert failed[0]["degree_type_source_urls"] == ""

    sourced = _results()
    sourced.people[0].result.profile.fields[ProfileField.full_name].sources = ["=unsafe-cell"]
    rich = list(
        csv.DictReader(
            io.StringIO(export_results(sourced, ["name"], "csv", provenance="field").decode("utf-8-sig"))
        )
    )
    assert rich[0]["full_name_source_urls"] == "'=unsafe-cell"


def test_export_rejects_unknown_provenance_mode():
    with pytest.raises(ValueError, match="Unsupported provenance mode"):
        export_results(_results(), ["name"], "csv", provenance="claims")


def test_export_has_each_enrichment_once_and_keeps_input_columns_unchanged():
    columns = ["name", "status", "status (2)", "full_name"]
    rows = list(csv.reader(io.StringIO(export_results(_results(), columns, "csv").decode("utf-8-sig"))))
    assert len(rows[0]) == len(columns) + 20
    assert len(set(rows[0])) == len(rows[0])
    assert rows[0][:4] == columns
    assert "status (3)" in rows[0] and "full_name (2)" in rows[0]
    assert columns == ["name", "status", "status (2)", "full_name"]


@pytest.mark.parametrize("format", ["csv", "xlsx"])
def test_export_escapes_headers_and_whitespace_prefixed_formulas(format):
    results = _results()
    results.people[0].original_row = {" =header": ' \t=HYPERLINK("https://example.org")'}
    data = export_results(results, [" =header"], format)
    if format == "csv":
        rows = list(csv.reader(io.StringIO(data.decode("utf-8-sig"))))
    else:
        workbook = load_workbook(io.BytesIO(data), data_only=False)
        rows = list(workbook.active.values)
        workbook.close()
    assert rows[0][0] == "' =header"
    assert rows[1][0].startswith("' \t=")


@pytest.mark.parametrize("hidden_cell", ["Z2", "A10000"])
def test_xlsx_cannot_hide_excess_cells_with_false_dimensions(hidden_cell):
    workbook = Workbook()
    workbook.active.append(["name"])
    workbook.active.append(["Jane Doe"])
    workbook.active[hidden_cell] = "unexpected"
    original = io.BytesIO()
    workbook.save(original)
    workbook.close()
    modified = io.BytesIO()
    with zipfile.ZipFile(original) as source, zipfile.ZipFile(modified, "w") as target:
        for item in source.infolist():
            content = source.read(item.filename)
            if item.filename == "xl/worksheets/sheet1.xml":
                content = re.sub(rb'<dimension ref="[^"]+"', b'<dimension ref="A1:A2"', content)
            target.writestr(item.filename, content)
    with pytest.raises(BatchParseError, match="exceed configured"):
        parse_batch(modified.getvalue(), "people.xlsx", None, _settings())


def test_xlsx_replaces_xml_invalid_input_without_losing_other_text():
    results = _results()
    results.people[0].original_row["name"] = "Jane\x00 Doe"
    data = export_results(results, ["name"], "xlsx")
    workbook = load_workbook(io.BytesIO(data))
    assert workbook.active.cell(2, 1).value == "Jane\ufffd Doe"
    workbook.close()
