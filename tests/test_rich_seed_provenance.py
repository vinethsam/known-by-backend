"""Supplied context remains an audit trail, never an extra evidence source."""

import csv
import io
import json
from datetime import date

import pytest
from openpyxl import load_workbook

from app.config import ScoringPolicy
from app.export.formats import export_results
from app.export.schema import standard_columns
from app.research.identity import assess_identity
from app.research.normalisation import normalise
from app.research.reconciliation import reconcile
from app.research.seed_provenance import attach_seed_provenance, supplied_attributes
from app.schemas import (
    EvidenceClaim,
    FieldDecision,
    JobResults,
    PersonResultView,
    PersonSeed,
    ProfileField,
    ResearchResult,
    SourceRecord,
    SourceType,
)

POLICY = ScoringPolicy()
TODAY = date(2026, 10, 23)
NAME = "Maya Silva"


def evidence_source(seed, text, identifier="official"):
    url = f"https://{identifier}.example/biography"
    return SourceRecord(
        source_id=identifier,
        person_id="person",
        requested_url=url,
        final_url=url,
        canonical_url=url,
        domain=f"{identifier}.example",
        source_type=SourceType.university,
        authority_score=0.95,
        identity=assess_identity(seed, text, POLICY, model_relevance="likely"),
        content_hash=identifier,
    )


def evidence_claim(identifier, field, value, *, quote=None, group=None, **kwargs):
    normalized, certainty = normalise(field, value)
    return EvidenceClaim(
        claim_id=identifier,
        person_id="person",
        source_id="official",
        field=field,
        raw_value=value,
        normalised_value=normalized,
        normalisation_certainty=certainty,
        evidence_text=quote or f"{NAME}: {value}.",
        subject_name=NAME,
        identity_relevance=1,
        extraction_model="offline-fixture",
        fact_group=group,
        **kwargs,
    )


def profile(seed, claims=(), sources=()):
    return reconcile("person", seed, list(claims), list(sources), POLICY, today=TODAY)


def rich_seed():
    return PersonSeed(
        full_name=NAME,
        country_of_origin="  Sri Lanka  ",
        organisation="Aster Labs",
        job_title="Research Director",
        university_name="Lake University",
        degree_type="MSc",
        subject="Physics",
        profile_link="https://official.example/biography",
        known_attributes={"LIST_NAME": "Private List", "Email": "private@example.test"},
    )


def test_all_supplied_core_fields_have_input_certainty_without_research_coverage():
    seed = rich_seed()
    result = profile(seed)

    assert result.profile_confidence == result.coverage == result.legacy_coverage == 0
    assert result.sources_used == 0
    assert len(result.seed_attributes) == 10
    assert all(item["input_confidence"] == 100 for item in result.seed_attributes)
    for field in ProfileField:
        decision = result.fields[field]
        assert decision.input_confidence == 100
        assert decision.research_confidence == 0
        assert decision.provenance_status == "supplied"
        assert decision.supplied_values
        assert not decision.supporting_source_ids
        assert not decision.supporting_claim_ids
        if field == ProfileField.full_name:
            assert decision.value == NAME
            assert decision.confidence == 100
        else:
            assert decision.value is None
            assert decision.confidence == 0


def test_an_alternate_researched_name_does_not_corroborate_immutable_seed_spelling():
    seed = PersonSeed(full_name=NAME)
    source = evidence_source(seed, f"{NAME} publishes under the name Maya A Silva.")
    alternate = evidence_claim("alternate-name", ProfileField.full_name, "Maya A Silva")
    result = profile(seed, [alternate], [source])
    decision = result.fields[ProfileField.full_name]

    assert decision.value == NAME
    assert decision.confidence == decision.input_confidence == 100
    assert decision.research_confidence == 0
    assert decision.provenance_status == "supplied"
    assert decision.supplied_values == [NAME]
    assert not decision.supporting_claim_ids
    assert decision.scoring_components["seed_corroborating_claim_ids"] == []


def test_attaching_seed_audit_does_not_mutate_independent_support_or_confidence():
    seed = PersonSeed(full_name=NAME, university_name="Lake University")
    source = evidence_source(seed, f"{NAME} graduated from Lake University.")
    claim = evidence_claim("university", ProfileField.university_name, "Lake University")
    decision = FieldDecision(
        value="Lake University",
        confidence=73,
        selected_claim_id=claim.claim_id,
        supporting_claim_ids=[claim.claim_id],
        supporting_source_ids=[source.source_id],
        scoring_components={"independent_sources": 1},
    )
    fields = {ProfileField.university_name: decision}
    attributes = supplied_attributes(seed, [claim], {source.source_id: source})

    attach_seed_provenance(fields, seed, attributes, FieldDecision())

    assert decision.confidence == decision.research_confidence == 73
    assert decision.input_confidence == 100
    assert decision.supporting_claim_ids == ["university"]
    assert decision.supporting_source_ids == ["official"]
    assert decision.scoring_components["independent_sources"] == 1
    assert decision.provenance_status == "corroborated"
    assert decision.scoring_components["seed_corroborating_claim_ids"] == ["university"]
    university_input = next(item for item in attributes if item["field"] == "university_name")
    assert university_input["source_urls"] == [source.canonical_url]


def test_same_external_evidence_has_same_confidence_with_corroborating_extra_seed_fields():
    simple = PersonSeed(full_name=NAME, university_name="Lake University")
    enriched = PersonSeed(
        full_name=NAME, university_name="Lake University", degree_type="MSc", subject="Physics"
    )
    source = evidence_source(simple, f"{NAME} earned an MSc in Physics from Lake University.")
    claims = [
        evidence_claim("university", ProfileField.university_name, "Lake University", group="masters"),
        evidence_claim("degree", ProfileField.degree_type, "MSc", group="masters"),
        evidence_claim("subject", ProfileField.subject, "Physics", group="masters"),
    ]

    baseline = profile(simple, claims, [source])
    richer = profile(enriched, claims, [source])

    assert richer.coverage == baseline.coverage
    assert richer.profile_confidence == baseline.profile_confidence
    assert richer.sources_used == baseline.sources_used == 1
    for field in ProfileField:
        assert richer.fields[field].confidence == baseline.fields[field].confidence
        assert richer.fields[field].supporting_claim_ids == baseline.fields[field].supporting_claim_ids
        assert richer.fields[field].supporting_source_ids == baseline.fields[field].supporting_source_ids
    assert richer.fields[ProfileField.degree_type].provenance_status == "corroborated"


def test_conflicting_country_preserves_exact_seed_and_external_evidence_separately():
    seed = PersonSeed(full_name=NAME, university_name="Lake University", country_of_origin="  Sri Lanka  ")
    text = f"{NAME} graduated from Lake University. {NAME} is from Canada."
    source = evidence_source(seed, text)
    claim = evidence_claim(
        "origin", ProfileField.country_of_origin, "Canada", quote=f"{NAME} is from Canada."
    )
    result = profile(seed, [claim], [source])
    decision = result.fields[ProfileField.country_of_origin]

    assert decision.value == "Canada"
    assert decision.supplied_values == ["  Sri Lanka  "]
    assert decision.provenance_status == "conflicting"
    assert decision.input_confidence == 100
    assert decision.research_confidence == decision.confidence > 0
    assert decision.supporting_claim_ids == ["origin"]
    assert decision.scoring_components["seed_conflicting_claim_ids"] == ["origin"]
    supplied = next(item for item in result.seed_attributes if item["field"] == "country_of_origin")
    assert supplied["value"] == "  Sri Lanka  "
    assert supplied["conflicting_claim_ids"] == ["origin"]
    assert supplied["source_urls"] == [source.canonical_url]


def test_old_seed_employer_is_superseded_when_evidence_supports_a_current_career_change():
    seed = PersonSeed(full_name=NAME, organisation="Former Labs")
    source = evidence_source(
        seed, f"{NAME} formerly worked for Former Labs. {NAME} now works for Aster Labs as Director."
    )
    claims = [
        evidence_claim(
            "former",
            ProfileField.organisation,
            "Former Labs",
            group="old",
            is_current=False,
            end_date=date(2020, 1, 1),
        ),
        evidence_claim("current", ProfileField.organisation, "Aster Labs", group="current", is_current=True),
        evidence_claim("role", ProfileField.job_title, "Director", group="current", is_current=True),
    ]
    result = profile(seed, claims, [source])
    decision = result.fields[ProfileField.organisation]

    assert decision.value == "Aster Labs"
    assert decision.selected_claim_id == "current"
    assert decision.supplied_values == ["Former Labs"]
    assert decision.input_confidence == 100
    assert decision.provenance_status == "superseded"
    supplied = next(item for item in result.seed_attributes if item["field"] == "organisation")
    assert supplied["provenance_status"] == "superseded"
    assert supplied["corroborating_claim_ids"] == ["former"]


def test_explicit_historical_seed_values_do_not_fill_current_employment():
    seed = PersonSeed(
        full_name=NAME, historical_employer="Former Labs", historical_job_title="Research Assistant"
    )
    result = profile(seed)

    assert result.fields[ProfileField.organisation].value is None
    assert result.fields[ProfileField.job_title].value is None
    assert result.fields[ProfileField.organisation].supplied_values == []
    assert result.fields[ProfileField.job_title].supplied_values == []
    assert result.coverage == 0
    historical = [item for item in result.seed_attributes if item["field"].startswith("historical_")]
    assert {item["value"] for item in historical} == {"Former Labs", "Research Assistant"}
    assert all(item["input_confidence"] == 100 for item in historical)


def test_explicit_historical_seed_keeps_corroborating_provenance_when_current_role_changes():
    seed = PersonSeed(full_name=NAME, historical_employer="Former Labs")
    source = evidence_source(
        seed, f"{NAME} formerly worked for Former Labs. {NAME} now works for Aster Labs as Director."
    )
    claims = [
        evidence_claim(
            "former",
            ProfileField.organisation,
            "Former Labs",
            group="old",
            is_current=False,
            end_date=date(2020, 1, 1),
        ),
        evidence_claim("current", ProfileField.organisation, "Aster Labs", group="current", is_current=True),
        evidence_claim("role", ProfileField.job_title, "Director", group="current", is_current=True),
    ]
    result = profile(seed, claims, [source])

    assert result.fields[ProfileField.organisation].value == "Aster Labs"
    assert result.fields[ProfileField.organisation].supplied_values == []
    historical = next(item for item in result.seed_attributes if item["field"] == "historical_employer")
    assert historical["value"] == "Former Labs"
    assert historical["input_confidence"] == 100
    assert historical["provenance_status"] == "superseded"
    assert historical["corroborating_claim_ids"] == ["former"]
    assert historical["source_urls"] == [source.canonical_url]


def test_multiple_education_records_do_not_borrow_unverified_seed_components():
    seed = PersonSeed(
        full_name=NAME,
        organisation="Aster Labs",
        university_name="Lake University",
        degree_type="MSc",
        subject="Physics",
    )
    source = evidence_source(
        seed, f"{NAME} works for Aster Labs. {NAME} graduated from Lake University and Hill University."
    )
    claims = [
        evidence_claim("lake", ProfileField.university_name, "Lake University", group="bachelor"),
        evidence_claim("bachelor", ProfileField.degree_type, "BSc", group="bachelor"),
        evidence_claim("hill", ProfileField.university_name, "Hill University", group="doctorate"),
        evidence_claim("doctorate", ProfileField.degree_type, "PhD", group="doctorate"),
        evidence_claim("chemistry", ProfileField.subject, "Chemistry", group="doctorate"),
    ]
    result = profile(seed, claims, [source])
    records = {record.fields[ProfileField.university_name].value: record for record in result.records}

    assert set(records) == {"Lake University", "Hill University"}
    lake = records["Lake University"].fields
    hill = records["Hill University"].fields
    assert lake[ProfileField.subject].value is None
    assert lake[ProfileField.subject].research_confidence == 0
    assert lake[ProfileField.subject].supplied_values == ["Physics"]
    assert hill[ProfileField.subject].value == "Chemistry"
    assert hill[ProfileField.subject].provenance_status == "researched"
    assert hill[ProfileField.university_name].provenance_status == "researched"
    assert lake[ProfileField.degree_type].value != "Master's Degree"
    assert hill[ProfileField.degree_type].value != "Master's Degree"
    for record in result.records:
        assert record.fields[ProfileField.full_name].value == NAME
        assert record.fields[ProfileField.organisation].value is None
        assert record.fields[ProfileField.organisation].supplied_values == ["Aster Labs"]
        assert record.fields[ProfileField.organisation].input_confidence == 100


@pytest.mark.parametrize("format", ["csv", "xlsx"])
def test_optional_export_provenance_keeps_fixed_prefix_and_raw_seed_certainty(format):
    seed = rich_seed()
    result = profile(seed)
    results = JobResults(
        job_id="job",
        status="completed",
        people=[
            PersonResultView(
                person_id="person",
                row_index=2,
                original_row=seed.input_metadata,
                status="review_required",
                result=ResearchResult(profile=result),
            )
        ],
    )
    output = export_results(results, list(seed.input_metadata), format, provenance="field")
    if format == "csv":
        reader = csv.DictReader(io.StringIO(output.decode("utf-8-sig")))
        headers = reader.fieldnames
        row = next(reader)
    else:
        workbook = load_workbook(io.BytesIO(output), read_only=True)
        records = workbook.active.iter_rows(values_only=True)
        headers = list(next(records))
        row = dict(zip(headers, next(records), strict=True))
        workbook.close()

    assert headers[: len(standard_columns())] == list(standard_columns())
    assert float(row["university_name_input_confidence"]) == 100
    assert float(row["university_name_research_confidence"]) == 0
    assert float(row["university_name_confidence"]) == 0
    assert row["university_name"] in {None, ""}
    assert json.loads(row["country_of_origin_supplied_values"]) == ["  Sri Lanka  "]
    metadata = {item["column"]: item for item in json.loads(row["seed_input_metadata"])}
    assert metadata["country_of_origin"]["value"] == "  Sri Lanka  "
    assert metadata["Email"]["value"] == "private@example.test"
    assert all(item["input_confidence"] == 100 for item in metadata.values())


@pytest.mark.parametrize("format", ["csv", "xlsx"])
def test_large_seed_provenance_round_trips_without_excel_cell_truncation(format):
    original = {"first unknown": "A" * 20000, "second unknown": "=-+@" * 5000}
    seed = PersonSeed(full_name=NAME, input_metadata=original)
    result = profile(seed)
    results = JobResults(
        job_id="job",
        status="completed",
        people=[
            PersonResultView(
                person_id="person",
                row_index=2,
                original_row=seed.input_metadata,
                status="review_required",
                result=ResearchResult(profile=result),
            )
        ],
    )
    output = export_results(results, list(seed.input_metadata), format, provenance="field")
    if format == "csv":
        reader = csv.DictReader(io.StringIO(output.decode("utf-8-sig")))
        headers = reader.fieldnames
        row = next(reader)
    else:
        workbook = load_workbook(io.BytesIO(output), read_only=True)
        records = workbook.active.iter_rows(values_only=True)
        headers = list(next(records))
        row = dict(zip(headers, next(records), strict=True))
        workbook.close()

    assert headers[: len(standard_columns())] == list(standard_columns())
    assert "seed_input_metadata__part_2" in headers
    fragments = [row["seed_input_metadata"]]
    index = 2
    while f"seed_input_metadata__part_{index}" in row:
        fragments.append(row[f"seed_input_metadata__part_{index}"] or "")
        index += 1
    assert all(len(fragment) <= 32767 for fragment in fragments)
    restored = {item["column"]: item for item in json.loads("".join(fragments))}
    for column, raw in original.items():
        assert restored[column]["value"] == raw
        assert restored[column]["input_confidence"] == 100


@pytest.mark.parametrize("format", ["csv", "xlsx"])
@pytest.mark.parametrize("status", ["failed", "cancelled"])
def test_unresearched_rows_preserve_input_certainty_and_long_metadata_without_research(format, status):
    original = {
        "Full Name": f"  {NAME}  ",
        "University": "  Lake University  ",
        "Previous Employer": "Former Labs",
        "Former Job Title": "Research Assistant",
        "private annotation": "A" * 20000,
        "other unknown": "=-+@" * 5000,
    }
    results = JobResults(
        job_id="job",
        status=status,
        people=[
            PersonResultView(
                person_id="person",
                row_index=2,
                original_row=original,
                status=status,
                error_code="PROVIDER_UNAVAILABLE" if status == "failed" else None,
            )
        ],
    )
    output = export_results(results, list(original), format, provenance="field")
    if format == "csv":
        reader = csv.DictReader(io.StringIO(output.decode("utf-8-sig")))
        headers = reader.fieldnames
        row = next(reader)
    else:
        workbook = load_workbook(io.BytesIO(output), read_only=True)
        records = workbook.active.iter_rows(values_only=True)
        headers = list(next(records))
        row = dict(zip(headers, next(records), strict=True))
        workbook.close()

    assert headers[: len(standard_columns())] == list(standard_columns())
    assert row["status"] == status
    assert row["error_code"] == (
        "PROVIDER_UNAVAILABLE" if status == "failed" else "" if format == "csv" else None
    )
    for field in ProfileField:
        assert row[field.value] in {"", None}
        assert row[f"{field.value}_confidence"] in {"", None}
        assert row[f"{field.value}_research_confidence"] in {"", None}
        assert row[f"{field.value}_source_urls"] in {"", None}
    assert row["profile_confidence"] in {"", None}
    assert row["coverage"] in {"", None}
    assert float(row["full_name_input_confidence"]) == 100
    assert float(row["university_name_input_confidence"]) == 100
    assert row["university_name_provenance_status"] == "supplied"
    assert json.loads(row["university_name_supplied_values"]) == ["  Lake University  "]
    assert json.loads(row["organisation_supplied_values"]) == []
    assert json.loads(row["job_title_supplied_values"]) == []
    assert row["organisation_input_confidence"] in {"", None}
    assert "seed_input_metadata__part_2" in row
    parts = [row["seed_input_metadata"]]
    index = 2
    while f"seed_input_metadata__part_{index}" in row:
        parts.append(row[f"seed_input_metadata__part_{index}"] or "")
        index += 1
    assert all(len(part) <= 32767 for part in parts)
    attributes = json.loads("".join(parts))
    assert {item["column"]: item["value"] for item in attributes} == original
    assert all(item["input_confidence"] == 100 for item in attributes)
    assert all(item["provenance_status"] == "supplied" for item in attributes)

    # Optional audit columns must not appear in the existing default export.
    default = export_results(results, list(original), "csv")
    default_reader = csv.DictReader(io.StringIO(default.decode("utf-8-sig")))
    assert "seed_input_metadata" not in default_reader.fieldnames
    assert next(default_reader)["University"] == original["University"]


@pytest.mark.parametrize(
    "text",
    [
        f"{NAME} never attended Lake University.",
        f"{NAME} did not graduate from Lake University.",
        f"{NAME} has no affiliation with Lake University.",
    ],
)
def test_explicit_university_denial_cannot_supply_an_identity_anchor(text):
    seed = PersonSeed(full_name=NAME, university_name="Lake University")
    identity = assess_identity(seed, text, POLICY)

    assert identity.signals["matched_seed_anchors"] == []
    assert identity.signals["contradicted_clues"] == ["university_name"]
    assert identity.score <= POLICY.identity_name_only
    assert identity.ambiguous
    assert "SEED_IDENTITY_CONTRADICTION" in identity.reason_codes
