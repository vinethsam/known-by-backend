"""Bad individual input records must never discard or research valid neighbours."""

import csv
import io

import pytest
from alembic.config import Config
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook
from sqlalchemy import select

from alembic import command
from app.api.routes import create_app
from app.config import Settings
from app.db.models import PersonTaskRow
from app.db.store import Store
from app.export.batch import parse_batch
from app.export.formats import export_results
from app.research.reconciliation import reconcile
from app.schemas import ResearchResult


def settings(tmp_path):
    return Settings(
        _env_file=None,
        APP_ENV="test",
        DATABASE_URL=f"sqlite:///{tmp_path / 'batch-input.db'}",
        OPENROUTER_API_KEY="offline-key",
        OPENROUTER_SOURCE_MODEL="offline/source",
        OPENROUTER_EXTRACTION_MODEL="offline/extraction",
        STATIC_FETCH_WORKER_URL="https://worker.example.org/fetch",
        STATIC_FETCH_WORKER_SECRET="offline-secret",
        PLAYWRIGHT_ENABLED=False,
    )


def migrated_store(config):
    migration = Config("alembic.ini")
    migration.set_main_option("sqlalchemy.url", config.DATABASE_URL)
    command.upgrade(migration, "head")
    return Store(config)


def spreadsheet(rows, format):
    if format == "csv":
        output = io.StringIO(newline="")
        csv.writer(output).writerows(rows)
        return output.getvalue().encode()
    workbook = Workbook()
    for row in rows:
        workbook.active.append(row)
    output = io.BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


@pytest.mark.parametrize("format", ["csv", "xlsx"])
def test_mixed_input_keeps_row_errors_and_only_valid_people_are_researched(tmp_path, format):
    config = settings(tmp_path)
    rows = [
        ["Full Name", "University", "Notes"],
        ["José 王", "Lake University", "first\nsecond"],
        ["", "Hill University", "missing name"],
        ["A" * 201, "Hill University", "name exceeds the identity limit"],
        ["Amara Silva", "Lake University", "valid neighbour"],
    ]
    batch = parse_batch(spreadsheet(rows, format), f"people.{format}", None, config)
    assert batch.row_errors == [None, "INVALID_PERSON_NAME", "INVALID_PERSON_SEED", None]
    assert batch.seeds[1:3] == [None, None]
    store = migrated_store(config)
    try:
        created = store.create_job(batch.seeds, batch.rows, batch.columns, row_errors=batch.row_errors)
        assert created.total_people == 4
        initial = store.get_results(created.job_id)
        assert [person.row_index for person in initial.people] == [2, 3, 4, 5]
        assert [person.status for person in initial.people] == ["queued", "failed", "failed", "queued"]
        assert [person.error_code for person in initial.people] == batch.row_errors

        claimed = []
        while lease := store.claim_task("offline-test"):
            claimed.append(lease.seed.full_name)
            result = ResearchResult(profile=reconcile(lease.person_id, lease.seed, [], [], config.SCORING))
            assert store.finish_task(lease, result)
        assert claimed == ["José 王", "Amara Silva"]
        assert store.get_job(created.job_id).status == "partial"
        with store.session_factory() as session:
            failed = session.scalars(select(PersonTaskRow).where(PersonTaskRow.status == "failed")).all()
            assert len(failed) == 2
            assert all(task.attempts == 0 and task.seed_json == {} for task in failed)
            assert all(task.completed_at is not None for task in failed)

        results = store.get_results(created.job_id)
        output = export_results(results, batch.columns, format, provenance="field")
        if format == "csv":
            exported = list(csv.DictReader(io.StringIO(output.decode("utf-8-sig"))))
        else:
            workbook = load_workbook(io.BytesIO(output), read_only=True)
            values = workbook.active.iter_rows(values_only=True)
            headers = next(values)
            exported = [dict(zip(headers, row, strict=True)) for row in values]
            workbook.close()
        assert len(exported) == 4
        assert exported[1]["error_code"] == "INVALID_PERSON_NAME"
        assert exported[2]["error_code"] == "INVALID_PERSON_SEED"
        assert exported[1]["University"] == "Hill University"
        assert float(exported[1]["university_name_input_confidence"]) == 100
        assert exported[1]["university_name"] in {"", None}
    finally:
        store.engine.dispose()


@pytest.mark.parametrize("name_column", [None, "Imported label"])
def test_all_invalid_names_remain_individual_failed_records(tmp_path, name_column):
    config = settings(tmp_path)
    header = name_column or "Full Name"
    content = spreadsheet([[header], ["123"], [""], ["\x00"]], "csv")
    batch = parse_batch(content, "people.csv", name_column, config)
    assert batch.seeds == [None, None, None]
    assert batch.row_errors == ["INVALID_PERSON_NAME"] * 3
    store = migrated_store(config)
    try:
        created = store.create_job(batch.seeds, batch.rows, batch.columns, row_errors=batch.row_errors)
        assert created.status == "failed"
        assert store.get_job(created.job_id).counts["failed"] == 3
        assert store.claim_task("offline-test") is None
    finally:
        store.engine.dispose()


def test_invalid_unicode_scalar_is_confined_to_its_record(tmp_path):
    config = settings(tmp_path)
    content = spreadsheet([["name", "note"], ["Jane Doe", "bad\uffffvalue"], ["Amara Silva", "valid"]], "csv")
    batch = parse_batch(content, "people.csv", None, config)
    assert batch.row_errors == ["INVALID_PERSON_SEED", None]
    assert batch.rows[0]["note"] == "bad\uffffvalue"
    assert batch.seeds[1].full_name == "Amara Silva"


def test_api_accepts_sanitized_records_and_reports_invalid_rows_without_research(tmp_path):
    config = settings(tmp_path)
    store = migrated_store(config)
    with TestClient(create_app(config, store=store)) as client:
        response = client.post(
            "/v1/research/batch",
            files={"file": ("people.csv", b"name,note\nJane\x00 Doe,first\vsecond\n,missing\n", "text/csv")},
        )
        assert response.status_code == 202
        created = response.json()
        assert created["total_people"] == 2
        result = client.get(f"/v1/jobs/{created['job_id']}/results").json()
        assert [person["status"] for person in result["people"]] == ["queued", "failed"]
        assert result["people"][0]["original_row"]["name"] == "Jane\x00 Doe"
        assert result["people"][1]["error_code"] == "INVALID_PERSON_NAME"
        assert all(person["result"] is None for person in result["people"])
        lease = store.claim_task("offline-test")
        assert lease.seed.full_name == "Jane Doe"
        assert store.claim_task("offline-test") is None


def test_store_requires_error_for_missing_seed(tmp_path):
    store = Store(settings(tmp_path))
    try:
        with pytest.raises(ValueError, match="matching row error"):
            store.create_job([None], [{}], ["name"])
        with pytest.raises(ValueError, match="matching row error"):
            store.create_job([None], [{}], ["name"], row_errors=[])
    finally:
        store.engine.dispose()
