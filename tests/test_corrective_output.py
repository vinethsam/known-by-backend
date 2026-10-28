"""Offline acceptance cases for supplied values and the standardized artifact."""

import csv
import io
import json
from types import SimpleNamespace

import pytest
from openpyxl import load_workbook

from app.config import ScoringPolicy, Settings
from app.db.store import Store
from app.export.batch import parse_batch
from app.export.formats import ExportCellLimitError, export_results
from app.export.schema import RESEARCH_FIELDS, standard_columns
from app.research.confidence import profile_scores
from app.research.normalisation import normalise_display
from app.research.reconciliation import reconcile
from app.research.seed_provenance import project_stored_seed
from app.schemas import (
    FieldDecision,
    JobResults,
    PersonResultView,
    PersonSeed,
    PersonStatus,
    ProfileField,
    ResearchResult,
)


def result_for(seed):
    profile = reconcile("person", seed, [], [], ScoringPolicy())
    return JobResults(
        job_id="job",
        status="failed",
        people=[
            PersonResultView(
                person_id="person",
                row_index=2,
                original_row=seed.input_metadata,
                status="failed",
                error_code="IDENTITY_UNRESOLVED",
                result=ResearchResult(profile=profile),
            )
        ],
    )


def test_all_eight_prefixed_columns_project_after_failed_research_without_evidence():
    values = {
        "ALUMNI_FULL_NAME": "john michael smith",
        "ALUMNI_COUNTRY_OF_ORIGIN": "ghana",
        "ALUMNI_ORGANISATION": "6a architects",
        "ALUMNI_JOB_TITLE_EXACT": "architect",
        "ALUMNI_UNIVERSITY_NAME": "university of ghana",
        "ALUMNI_DEGREE_TYPE": "BSc",
        "ALUMNI_SUBJECT": "architecture",
        "ALUMNI_PROFILE_LINK": "https://example.com/John",
        "LIST_NAME": "private cohort",
    }
    content = io.StringIO()
    writer = csv.writer(content)
    writer.writerow(values)
    writer.writerow(values.values())
    batch = parse_batch(content.getvalue().encode(), "input.csv", None, Settings(APP_ENV="test"))
    seed = batch.seeds[0]
    results = result_for(seed)
    profile = results.people[0].result.profile
    expected = [
        "John Michael Smith",
        "Ghana",
        "6a Architects",
        "Architect",
        "University of Ghana",
        "Bachelor's Degree",
        "Architecture",
        "https://example.com/John",
    ]
    assert [profile.fields[field].value for field in RESEARCH_FIELDS] == expected
    for decision in profile.fields.values():
        assert decision.input_confidence == 100
        assert decision.research_confidence == 0
        assert decision.researched_value is None
        assert decision.provenance_status == "supplied"
        assert not decision.selected_claim_id and not decision.supporting_source_ids
    assert profile.coverage == profile.legacy_coverage == profile.profile_confidence == 0
    assert profile_scores(profile.fields, ScoringPolicy()) == (0, 0)
    repeated = project_stored_seed(project_stored_seed(profile, seed), seed)
    assert [repeated.fields[field].value for field in RESEARCH_FIELDS] == expected
    assert all(decision.provenance_status == "supplied" for decision in repeated.fields.values())
    assert seed.input_metadata == values


def test_failed_before_profile_uses_persisted_custom_mapping_and_keeps_raw_values():
    seed = PersonSeed(
        full_name="john smith",
        input_metadata={"who": "john smith", "work": "IBM"},
        input_field_map={"who": "full_name", "work": "organisation"},
    )
    task = SimpleNamespace(
        person_id="person",
        row_index=2,
        original_row_json={"who": "john smith", "work": "IBM"},
        seed_json=seed.model_dump(mode="json"),
        status="failed",
        error_code="PROVIDER_UNAVAILABLE",
    )
    view = Store._person_result_view(object.__new__(Store), task, None, [], [], [])
    assert view.result is None
    assert view.supplied_fields[ProfileField.organisation].value == "IBM"
    assert view.supplied_fields[ProfileField.organisation].input_confidence == 100
    assert view.supplied_fields[ProfileField.organisation].research_confidence == 0
    assert view.seed_attributes[1]["value"] == "IBM"


@pytest.mark.parametrize("format", ["csv", "xlsx"])
def test_excluded_supplied_profile_url_stays_in_details_without_becoming_a_finding(format):
    seed = PersonSeed(full_name="John Smith", profile_link="https://linkedin.com/in/john")
    results = result_for(seed)
    profile = results.people[0].result.profile
    for projected in (profile, project_stored_seed(profile, seed)):
        decision = projected.fields[ProfileField.profile_link]
        assert decision.value is None
        assert decision.researched_value is None
        assert decision.supplied_values == [seed.profile_link]
        assert decision.input_confidence == 100
        assert decision.research_confidence == 0
        assert decision.sources == []
        assert decision.scoring_components["seed_projection_reason"] == "EXCLUDED_SOURCE_POLICY"
    data = export_results(results, [], format)
    if format == "csv":
        row = next(csv.DictReader(io.StringIO(data.decode("utf-8-sig"))))
    else:
        workbook = load_workbook(io.BytesIO(data))
        rows = iter(workbook.active.values)
        row = dict(zip(next(rows), next(rows), strict=True))
        workbook.close()
    assert row["profile_link"] in {None, ""}
    assert float(row["profile_link_confidence"]) == 0
    assert float(row["profile_link_input_confidence"]) == 100
    assert json.loads(row["profile_link_supplied_values"]) == [seed.profile_link]


@pytest.mark.parametrize(
    "label", ["Nationality", "Citizenship", "Country of Birth", "Residence Country", "Employment Location"]
)
def test_geographic_seed_is_retained_without_becoming_origin(label):
    seed = PersonSeed(full_name="John Smith", input_metadata={"Full Name": "John Smith", label: "Ghana"})
    profile = result_for(seed).people[0].result.profile
    assert profile.fields[ProfileField.country_of_origin].value is None
    assert next(item for item in profile.seed_attributes if item["column"] == label)["value"] == "Ghana"


@pytest.mark.parametrize(
    "raw,expected",
    [
        ("ibm", "IBM"),
        ("unesco", "UNESCO"),
        ("iccrom", "ICCROM"),
        ("mcgill", "McGill"),
        ("pwc", "PwC"),
        ("eth zürich", "ETH Zürich"),
        ("rpbw", "RPBW"),
    ],
)
def test_professional_brand_casing(raw, expected):
    assert normalise_display(ProfileField.organisation, raw) == expected


def test_export_format_equivalence_review_fill_and_person_record_status():
    results = result_for(PersonSeed(full_name="John Smith", organisation="Old Practice"))
    person = results.people[0]
    person.status = PersonStatus.review_required
    fields = person.result.profile.records[0].fields
    fields[ProfileField.organisation] = FieldDecision(
        value="New Practice",
        confidence=92,
        research_confidence=92,
        supplied_values=["Old Practice"],
        input_confidence=100,
        provenance_status="conflicting",
        review_required=True,
        review_reason_codes=["SOURCE_CONFLICT"],
    )
    fields[ProfileField.job_title] = FieldDecision(
        value="Architect", confidence=49, review_required=True, review_reason_codes=["LOW_CONFIDENCE"]
    )
    fields[ProfileField.subject] = FieldDecision(value="Mathematics", confidence=90, review_required=False)
    first = person.result.profile.records[0]
    first.status = PersonStatus.review_required
    first.review_required = True
    second = first.model_copy(deep=True)
    second.status = PersonStatus.completed
    person.result.profile.records.append(second)
    rows = list(
        csv.DictReader(io.StringIO(export_results(results, ["LIST_NAME"], "csv").decode("utf-8-sig")))
    )
    workbook = load_workbook(io.BytesIO(export_results(results, ["LIST_NAME"], "xlsx")))
    sheet = workbook.active
    headers = [cell.value for cell in sheet[1]]
    assert headers == list(standard_columns()) and len(headers) == 73
    assert headers[:8] == [field.value for field in RESEARCH_FIELDS] and headers[-1] == "source_urls"
    assert "LIST_NAME" not in headers
    assert {row["status"] for row in rows} == {"review_required"}
    # Both records share the same conflicting employment context, so an old
    # completed record cannot hide the newly explicit supplied-value conflict.
    assert [row["record_status"] for row in rows] == ["review_required", "review_required"]

    def cell(key):
        return sheet.cell(2, headers.index(key) + 1)

    for key in ("organisation", "organisation_confidence", "job_title", "review_required"):
        assert cell(key).fill.fgColor.rgb.endswith("FFF2CC")
    for key in ("subject", "subject_confidence", "university_name", "full_name"):
        assert cell(key).fill.fill_type is None
    for key in headers:
        value = cell(key).value
        if isinstance(value, (int, float)) and not isinstance(value, bool):
            assert float(rows[0][key]) == value
        else:
            assert str(value if value is not None else "") == rows[0][key]
    assert json.loads(rows[0]["organisation_supplied_values"]) == ["Old Practice"]
    assert rows[0]["organisation_review_reasons"] == "SOURCE_CONFLICT; SUPPLIED_VALUE_CONFLICT"
    workbook.close()


@pytest.mark.parametrize("format", ["csv", "xlsx"])
def test_oversize_standard_provenance_fails_explicitly_without_losing_original(format):
    raw = "=" * 20000
    seed = PersonSeed(full_name="John Smith", input_metadata={"Organisation": raw})
    results = result_for(seed)
    with pytest.raises(ExportCellLimitError, match="32767"):
        export_results(results, [], format)
    assert results.people[0].result.profile.seed_attributes[0]["value"] == raw
